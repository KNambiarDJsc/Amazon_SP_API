import io

import pytest
from openpyxl import load_workbook
from support import brands_of, submit

from app import cli
from app.auth import authenticate
from app.config import get_settings
from app.db import session_scope


def test_demo_runs_offline_and_writes_a_labelled_workbook(engine, tmp_path, capsys):
    out = tmp_path / "report.xlsx"

    assert cli.main(["demo", "--out", str(out), "--database-url", "sqlite://"]) == 0

    printed = capsys.readouterr().out
    assert "DEMO DATA" in printed
    assert "Havells      BRAND_APPROVAL_REQUIRED" in printed
    assert "Oddy         NO_PRODUCT_FOUND" in printed
    wb = load_workbook(io.BytesIO(out.read_bytes()))
    assert wb["Verified Listable Products"]["A1"].value.startswith("DEMO DATA")
    assert wb["Verified Listable Products"].max_row == 2 + 6


def test_demo_accepts_custom_brands(engine, tmp_path, capsys):
    out = tmp_path / "r.xlsx"
    assert cli.main(["demo", "--out", str(out), "Havells", "Nobody"]) == 0
    printed = capsys.readouterr().out
    assert "Havells" in printed and "Nobody       NO_PRODUCT_FOUND" in printed


def test_create_seller_prints_a_working_key_exactly_once(engine, capsys):
    assert cli.main(["create-seller", "Acme Traders", "--amazon-seller-id", "A1ACME"]) == 0
    printed = capsys.readouterr().out
    key = next(line.split(":", 1)[1].strip() for line in printed.splitlines() if "api key" in line)
    assert key.startswith("spk_") and len(key) > 30
    assert "cannot be shown again" in printed

    with session_scope() as s:
        account = authenticate(s, key)
        assert account is not None and account.amazon_seller_id == "A1ACME"
        assert key not in (account.api_key_hash, account.api_key_prefix)  # never stored
        assert account.api_key_prefix == key[:12]


def test_drain_refuses_unconfigured_live_mode_instead_of_faking_results(
    engine, monkeypatch, capsys
):
    monkeypatch.setenv("AMAZON_MODE", "live")
    monkeypatch.delenv("LWA_CLIENT_ID", raising=False)
    get_settings.cache_clear()
    try:
        assert cli.main(["drain"]) == 2
    finally:
        get_settings.cache_clear()
    assert "LWA_CLIENT_ID is not set" in capsys.readouterr().err


def test_drain_runs_open_jobs(engine, seller, clock, capsys, monkeypatch):
    job_id = submit(seller, ["Havells"], clock)
    monkeypatch.setenv("RATE_LIMITER", "memory")
    get_settings.cache_clear()
    try:
        assert cli.main(["drain"]) == 0
    finally:
        get_settings.cache_clear()
    assert "ran 1 brand(s)" in capsys.readouterr().out
    assert brands_of(job_id)["Havells"].brand_status.value == "BRAND_APPROVAL_REQUIRED"


def test_a_command_is_required():
    with pytest.raises(SystemExit):
        cli.main([])


# ---- seller authorisation ------------------------------------------------------------------


@pytest.fixture
def fake_live(monkeypatch):
    from fake_amazon import FakeAmazon, live_settings

    from app.amazon.factory import LiveServices

    fake = FakeAmazon()
    services = LiveServices(live_settings(), http=fake.http())
    monkeypatch.setattr(cli, "_live_services", lambda: services)
    return fake, services


def _stdin(monkeypatch, text):
    monkeypatch.setattr("sys.stdin", io.StringIO(text))


def test_generate_key_prints_a_usable_key(capsys):
    from app.crypto import CredentialBox

    assert cli.main(["generate-key"]) == 0
    key = capsys.readouterr().out.strip()
    CredentialBox([key])  # does not raise


def test_connect_seller_reads_the_token_from_stdin_and_never_prints_it(
    engine, seller, fake_live, monkeypatch, capsys
):
    from fake_amazon import REFRESH_TOKEN, SELLING_PARTNER_ID

    _stdin(monkeypatch, REFRESH_TOKEN + "\n")
    code = cli.main(["connect-seller", str(seller.id), "--selling-partner-id", SELLING_PARTNER_ID])
    printed = capsys.readouterr()
    assert code == 0, printed.err
    assert "ACTIVE" in printed.out and "SELF_AUTHORIZED" in printed.out
    assert REFRESH_TOKEN not in printed.out + printed.err


def test_connect_seller_refuses_a_token_amazon_rejects(
    engine, seller, fake_live, monkeypatch, capsys
):
    from fake_amazon import SELLING_PARTNER_ID

    fake, _ = fake_live
    fake.revoke()
    _stdin(monkeypatch, "Atzr|revoked\n")
    code = cli.main(["connect-seller", str(seller.id), "--selling-partner-id", SELLING_PARTNER_ID])
    assert code == 1
    assert "invalid_grant" in capsys.readouterr().err


def test_rotate_credentials_moves_tokens_to_the_new_key(
    engine, seller, fake_live, monkeypatch, capsys
):
    from fake_amazon import ENCRYPTION_KEY, REFRESH_TOKEN, SELLING_PARTNER_ID

    from app.crypto import CredentialBox, generate_key
    from app.models import SellerAuthorization

    _stdin(monkeypatch, REFRESH_TOKEN + "\n")
    cli.main(["connect-seller", str(seller.id), "--selling-partner-id", SELLING_PARTNER_ID])
    new_key = generate_key()
    monkeypatch.setenv("CREDENTIALS_ENCRYPTION_KEYS", f"{new_key},{ENCRYPTION_KEY}")
    get_settings.cache_clear()
    try:
        assert cli.main(["rotate-credentials"]) == 0
    finally:
        get_settings.cache_clear()
    assert "re-encrypted 1 credential(s)" in capsys.readouterr().out
    with session_scope() as s:
        sealed = s.query(SellerAuthorization).one().refresh_token_encrypted
    assert CredentialBox([new_key]).open(seller.id, sealed) == REFRESH_TOKEN


def test_seller_status_for_an_unconnected_seller(engine, seller, capsys):
    assert cli.main(["seller-status", str(seller.id)]) == 0
    assert "connected            False" in capsys.readouterr().out
