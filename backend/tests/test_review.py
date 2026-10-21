"""Human-assisted verification: the review queue, recording what was seen, recheck and retry."""

import io

import pytest
from apienv import Env
from openpyxl import load_workbook
from sqlalchemy import select

from app.amazon.errors import AmazonTransientError
from app.amazon.mock import (
    NO_RESTRICTIONS,
    brand_gated,
    catalog_item,
    demo_client,
)
from app.db import session_scope
from app.models import AuditEvent, CheckObservation

LIMITATION = (
    "Used, Collectible, Refurbished conditions: "
    "This product has other listing limitations. Learn more"
)


def two_unknown_asins():
    """Oddy has two ASINs the restrictions API does not recognise -> both need a human."""
    items = [
        catalog_item(f"B0ODDY000{i}", f"Oddy product {i}", "Oddy", category_id="c") for i in (1, 2)
    ]
    client = demo_client()
    client.catalog["oddy"] = items
    return client


@pytest.fixture
def env(engine):
    return Env(two_unknown_asins())


@pytest.fixture
def auth(env):
    return env.new_seller()[1]


def run(env, auth, brands=("Oddy", "Nobody", "Havells")):
    job = env.http.post("/api/v1/jobs", json={"brands": list(brands)}, headers=auth).json()["id"]
    env.drain()
    return job


def queue(env, auth, job):
    return env.http.get(f"/api/v1/jobs/{job}/review-queue", headers=auth).json()


def check_for(env, auth, job, asin):
    r = env.http.get(f"/api/v1/jobs/{job}/checks", params={"q": asin}, headers=auth).json()
    return r["items"][0]["id"]


def evidence(**kw):
    base = dict(
        page_asin="B0ODDY0001",
        selected_condition="new_new",
        action_label="Sell this product",
        action_enabled=True,
        notices=[LIMITATION],
    )
    return {**base, **kw}


def submit(env, auth, job, check, **kw):
    return env.http.post(
        f"/api/v1/jobs/{job}/checks/{check}/manual-evidence", json=evidence(**kw), headers=auth
    )


# ---- the queue -------------------------------------------------------------------------------


def test_queue_lists_checks_and_brands_a_person_must_look_at(env, auth):
    job = run(env, auth)
    q = queue(env, auth, job)
    kinds = [(i["kind"], i["brand"], i["asin"]) for i in q["items"]]
    assert kinds == [
        ("CHECK", "Oddy", "B0ODDY0001"),
        ("CHECK", "Oddy", "B0ODDY0002"),
        ("BRAND", "Oddy", None),  # every ASIN unresolved -> the brand is unresolved too
        ("BRAND", "Nobody", None),
    ]  # Havells is settled and absent
    check = q["items"][0]
    assert check["status"] == "NEEDS_MANUAL_CHECK" and check["rule_id"] == "SPAPI_ASIN_NOT_FOUND"
    assert check["actions"] == ["record_evidence", "recheck"] and check["condition"] == "new_new"
    assert "B0ODDY0001" in check["how_to"][0] and "'New'" in check["how_to"][1]
    assert "Do not click" in check["how_to"][1]
    assert check["seller_central_url"] == "https://sellercentral.amazon.in"
    brand = q["items"][3]
    assert brand["status"] == "NO_PRODUCT_FOUND" and brand["actions"] == ["retry_brand"]


def test_queue_is_empty_when_nothing_needs_a_person(env, auth):
    job = run(env, auth, ["Havells"])
    assert queue(env, auth, job)["total"] == 0


# ---- recording what was seen -----------------------------------------------------------------


def test_screen_evidence_resolves_a_check_the_api_could_not(env, auth):
    job = run(env, auth)
    cid = check_for(env, auth, job, "B0ODDY0001")
    r = submit(env, auth, job, cid, note="Checked in Seller Central at 14:05")
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["status"] == "VERIFIED_LISTABLE" and body["rule_id"] == "BROWSER_SELL_ENABLED"
    assert body["verification_methods"] == ["MANUAL"]  # attributed to a person, not the API
    sources = [o["source"] for o in body["observations"]]
    assert sources == ["SP_API_RESTRICTIONS", "MANUAL"]  # the API's answer is kept as history
    manual = body["observations"][1]
    assert manual["raw"]["note"] == "Checked in Seller Central at 14:05"
    assert manual["raw"]["observer"].startswith("api-key:spk_")
    assert manual["signals"]["action_label"] == "Sell this product"
    assert manual["context"]["brand_validated"] is True
    assert manual["evidence_ref"].startswith("db:check_observations/")

    q = queue(env, auth, job)
    assert [i["asin"] for i in q["items"] if i["kind"] == "CHECK"] == ["B0ODDY0002"]  # shrank
    with session_scope() as s:
        event = s.scalars(
            select(AuditEvent).where(AuditEvent.event_type == "MANUAL_EVIDENCE_ADDED")
        ).one()
    assert event.payload["previous_status"] == "NEEDS_MANUAL_CHECK"
    assert event.payload["new_status"] == "VERIFIED_LISTABLE"


def test_resolving_every_check_updates_the_brand_summary_and_the_queue(env, auth):
    job = run(env, auth)
    for asin in ("B0ODDY0001", "B0ODDY0002"):
        assert (
            submit(env, auth, job, check_for(env, auth, job, asin), page_asin=asin).status_code
            == 200
        )
    brand = env.http.get(f"/api/v1/jobs/{job}/brands", params={"q": "oddy"}, headers=auth).json()[
        "items"
    ][0]
    assert brand["brand_status"] == "VERIFIED_LISTABLE"
    assert brand["brand_summary"]["status_counts"] == {"VERIFIED_LISTABLE": 2}
    assert brand["brand_summary"]["brand_wide_claim"] is False
    assert [i["brand"] for i in queue(env, auth, job)["items"]] == ["Nobody"]


def test_a_gated_screen_is_recorded_as_a_restriction(env, auth):
    job = run(env, auth)
    cid = check_for(env, auth, job, "B0ODDY0001")
    r = submit(
        env,
        auth,
        job,
        cid,
        notices=["You need approval to list in this brand."],
        action_label="Apply to sell",
    )
    assert r.json()["status"] == "BRAND_APPROVAL_REQUIRED"
    assert r.json()["messages"] == ["You need approval to list in this brand."]


@pytest.mark.parametrize(
    ("overrides", "rule"),
    [
        (dict(selected_condition=None, action_enabled=False), "BROWSER_CONDITION_NOT_SELECTED"),
        (dict(page_asin="B0OTHER000"), "BROWSER_ASIN_MISMATCH"),
        (dict(action_enabled=False), "BROWSER_ACTION_DISABLED"),
        (dict(notices=["Something unexpected"]), "BROWSER_UNRECOGNIZED_NOTICE"),
        (dict(product_unavailable=True), "BROWSER_PRODUCT_UNAVAILABLE"),
        (dict(selected_condition="used_good"), "BROWSER_CONDITION_NOT_SELECTED"),
    ],
)
def test_screen_states_that_do_not_prove_listability_leave_the_check_unresolved(
    env, auth, overrides, rule
):
    job = run(env, auth)
    cid = check_for(env, auth, job, "B0ODDY0001")
    body = submit(env, auth, job, cid, **overrides).json()
    assert body["status"] == "NEEDS_MANUAL_CHECK"
    assert body["observations"][-1]["source"] == "MANUAL"
    # The reviewer is told why their entry was not enough (the latest evidence leads), and the
    # API's earlier reason is still there.
    assert body["rule_id"] == rule  # not the API's older SPAPI_ASIN_NOT_FOUND


def test_an_unresolved_check_explains_why_and_keeps_the_earlier_reason(env, auth):
    job = run(env, auth)
    cid = check_for(env, auth, job, "B0ODDY0001")
    body = submit(env, auth, job, cid, selected_condition=None, action_enabled=False).json()
    assert body["rule_id"] == "BROWSER_CONDITION_NOT_SELECTED"  # what to fix, first
    first, second = body["rationale"].split(" | ")
    assert "Condition new_new is not selected" in first
    assert "does not recognise it in this marketplace" in second  # the API's reason is kept


def test_a_reviewer_cannot_supply_a_status_or_any_unknown_field(env, auth):
    job = run(env, auth)
    cid = check_for(env, auth, job, "B0ODDY0001")
    for extra in (
        {"status": "VERIFIED_LISTABLE"},
        {"verdict": "ungated"},
        {"verification_methods": []},
    ):
        assert submit(env, auth, job, cid, **extra).status_code == 422


@pytest.mark.parametrize(
    "bad",
    [
        dict(page_asin="short"),
        dict(page_asin="B0ODDY0001X"),
        dict(page_asin="B0ODDY 001"),
        dict(selected_condition="brand_new"),
        dict(notices=["x"] * 11),
        dict(notices=["x" * 501]),
        dict(action_label="x" * 61),
        dict(note="n" * 501),
    ],
)
def test_malformed_evidence_is_422(env, auth, bad):
    job = run(env, auth)
    assert submit(env, auth, job, check_for(env, auth, job, "B0ODDY0001"), **bad).status_code == 422


def test_asin_is_normalised_to_upper_case(env, auth):
    job = run(env, auth)
    cid = check_for(env, auth, job, "B0ODDY0001")
    assert (
        submit(env, auth, job, cid, page_asin="b0oddy0001").json()["status"] == "VERIFIED_LISTABLE"
    )


def test_evidence_for_a_check_that_has_not_completed_is_refused(env, auth):
    import uuid

    from app.enums import CheckState
    from app.models import EligibilityCheck

    job = run(env, auth)
    cid = check_for(env, auth, job, "B0ODDY0001")
    with session_scope() as s:  # a check that is planned but whose API call has not finished
        check = s.get(EligibilityCheck, uuid.UUID(cid))
        check.state, check.status = CheckState.PENDING, None
    r = submit(env, auth, job, cid)
    assert r.status_code == 409 and "not completed" in r.json()["detail"]


def test_cross_seller_access_is_404(env, auth):
    job = run(env, auth)
    cid = check_for(env, auth, job, "B0ODDY0001")
    _, other = env.new_seller("Other")
    assert submit(env, other, job, cid).status_code == 404
    assert env.http.get(f"/api/v1/jobs/{job}/review-queue", headers=other).status_code == 404
    assert (
        env.http.post(f"/api/v1/jobs/{job}/checks/{cid}/recheck", headers=other).status_code == 404
    )


# ---- conflicts surface instead of being resolved by whoever spoke last -----------------------


def test_a_reviewer_cannot_override_an_api_that_says_gated(env, auth):
    env.amazon.restrictions["B0ODDY0001"] = brand_gated("B0ODDY0001")
    job = run(env, auth)
    cid = check_for(env, auth, job, "B0ODDY0001")
    assert env.http.get(f"/api/v1/jobs/{job}/checks/{cid}", headers=auth).json()["status"] == (
        "BRAND_APPROVAL_REQUIRED"
    )
    body = submit(env, auth, job, cid).json()  # reviewer says listable
    assert body["status"] == "NEEDS_MANUAL_CHECK" and body["rule_id"] == "SOURCE_CONFLICT"
    assert "SP_API_RESTRICTIONS=BRAND_APPROVAL_REQUIRED" in body["rationale"]
    queued = [i["check_id"] for i in queue(env, auth, job)["items"] if i["kind"] == "CHECK"]
    assert cid in queued and len(queued) == 2  # the conflict stays in the queue


def test_recheck_asks_the_api_again_and_combines_it_with_the_screen_evidence(env, auth):
    job = run(env, auth)
    cid = check_for(env, auth, job, "B0ODDY0001")
    other = check_for(env, auth, job, "B0ODDY0002")
    submit(env, auth, job, cid)  # reviewer: listable (API had not recognised the ASIN)
    env.amazon.restrictions["B0ODDY0001"] = NO_RESTRICTIONS  # Amazon's catalogue caught up
    calls_before = list(env.amazon.calls_to("getListingsRestrictions"))

    r = env.http.post(f"/api/v1/jobs/{job}/checks/{cid}/recheck", headers=auth)
    assert r.status_code == 200 and r.json()["state"] == "PENDING" and r.json()["status"] is None
    assert (
        env.http.get(f"/api/v1/jobs/{job}", headers=auth).json()["status"] == "RUNNING"
    )  # reopened
    env.drain()

    body = env.http.get(f"/api/v1/jobs/{job}/checks/{cid}", headers=auth).json()
    assert body["status"] == "VERIFIED_LISTABLE"
    assert set(body["verification_methods"]) == {"SP_API_RESTRICTIONS", "MANUAL"}  # now agreeing
    assert [o["source"] for o in body["observations"]] == [
        "SP_API_RESTRICTIONS",
        "MANUAL",
        "SP_API_RESTRICTIONS",
    ]
    new_calls = env.amazon.calls_to("getListingsRestrictions")[len(calls_before) :]
    assert [c["asin"] for c in new_calls] == ["B0ODDY0001"]  # only that ASIN was queried again
    assert env.http.get(f"/api/v1/jobs/{job}/checks/{other}", headers=auth).json()["status"] == (
        "NEEDS_MANUAL_CHECK"
    )
    assert env.http.get(f"/api/v1/jobs/{job}", headers=auth).json()["status"] == "COMPLETED"


def test_recheck_of_a_busy_or_unfinished_check_is_409(env, auth):
    job = run(env, auth)
    cid = check_for(env, auth, job, "B0ODDY0001")
    assert (
        env.http.post(f"/api/v1/jobs/{job}/checks/{cid}/recheck", headers=auth).status_code == 200
    )
    again = env.http.post(f"/api/v1/jobs/{job}/checks/{cid}/recheck", headers=auth)
    assert again.status_code == 409  # brand is already queued
    env.http.post(f"/api/v1/jobs/{job}/cancel", headers=auth)
    assert (
        env.http.post(f"/api/v1/jobs/{job}/checks/{cid}/recheck", headers=auth).status_code == 409
    )


# ---- retrying brands -------------------------------------------------------------------------


def test_a_dead_lettered_brand_is_retried_from_its_checkpoint(env, auth):
    env.amazon.catalog_behaviour = AmazonTransientError("search is down")
    job = run(env, auth, ["Havells"])
    brand = env.http.get(f"/api/v1/jobs/{job}/brands", headers=auth).json()["items"][0]
    assert brand["status"] == "FAILED" and brand["brand_status"] == "NEEDS_MANUAL_CHECK"
    item = queue(env, auth, job)["items"][0]
    assert (
        item["kind"] == "BRAND"
        and item["actions"] == ["retry_brand"]
        and "search is down" in item["rationale"]
    )

    env.amazon.catalog_behaviour = None  # the outage is over
    r = env.http.post(f"/api/v1/jobs/{job}/brands/{brand['id']}/retry", headers=auth)
    assert r.status_code == 200 and r.json()["status"] == "RUNNING"
    env.drain()
    done = env.http.get(f"/api/v1/jobs/{job}/brands", headers=auth).json()["items"][0]
    assert done["status"] == "COMPLETED" and done["brand_status"] == "BRAND_APPROVAL_REQUIRED"
    assert done["attempts"] == 0 and done["last_error"] is None
    assert queue(env, auth, job)["total"] == 0


def test_no_product_found_can_be_rediscovered(env, auth):
    job = run(env, auth, ["Nobody"])
    brand = env.http.get(f"/api/v1/jobs/{job}/brands", headers=auth).json()["items"][0]
    assert brand["brand_status"] == "NO_PRODUCT_FOUND"
    env.amazon.catalog["nobody"] = [catalog_item("B0NOBODY01", "Nobody thing", "Nobody")]
    env.amazon.restrictions["B0NOBODY01"] = NO_RESTRICTIONS
    r = env.http.post(
        f"/api/v1/jobs/{job}/brands/{brand['id']}/retry", json={"rediscover": True}, headers=auth
    )
    assert r.status_code == 200
    env.drain()
    done = env.http.get(f"/api/v1/jobs/{job}/brands", headers=auth).json()["items"][0]
    assert done["brand_status"] == "VERIFIED_LISTABLE"
    assert len(env.amazon.calls_to("searchCatalogItems")) == 2


def test_retry_validations(env, auth):
    job = run(env, auth, ["Havells"])
    brand = env.http.get(f"/api/v1/jobs/{job}/brands", headers=auth).json()["items"][0]["id"]
    url = f"/api/v1/jobs/{job}/brands/{brand}/retry"
    assert env.http.post(url, headers=auth).status_code == 200
    assert env.http.post(url, headers=auth).status_code == 409  # already queued
    env.drain()
    _, other = env.new_seller("Other")
    assert env.http.post(url, headers=other).status_code == 404
    import uuid

    assert (
        env.http.post(f"/api/v1/jobs/{job}/brands/{uuid.uuid4()}/retry", headers=auth).status_code
        == 404
    )
    # A completed job cannot be cancelled; a job cancelled while running refuses retries.
    assert env.http.post(f"/api/v1/jobs/{job}/cancel", headers=auth).status_code == 409
    running = env.http.post("/api/v1/jobs", json={"brands": ["Havells"]}, headers=auth).json()["id"]
    env.http.post(f"/api/v1/jobs/{running}/cancel", headers=auth)
    cancelled_brand = env.http.get(f"/api/v1/jobs/{running}/brands", headers=auth).json()["items"][
        0
    ]
    r = env.http.post(f"/api/v1/jobs/{running}/brands/{cancelled_brand['id']}/retry", headers=auth)
    assert r.status_code == 409 and "cancelled" in r.json()["detail"]


# ---- the export reflects the resolution ------------------------------------------------------


def test_resolved_items_move_from_manual_review_to_the_verified_sheet(env, auth):
    job = run(env, auth, ["Oddy"])
    cid = check_for(env, auth, job, "B0ODDY0001")
    submit(env, auth, job, cid)
    wb = load_workbook(
        io.BytesIO(env.http.get(f"/api/v1/jobs/{job}/export.xlsx", headers=auth).content)
    )

    def rows(name):
        ws = wb[name]
        header = [c.value for c in ws[1]]
        return [
            dict(zip(header, [c.value for c in r], strict=True)) for r in ws.iter_rows(min_row=2)
        ]

    verified = rows("Verified Listable Products")
    assert [(r["ASIN"], r["Verification method"]) for r in verified] == [("B0ODDY0001", "MANUAL")]
    assert [r["ASIN"] for r in rows("Manual Review") if r["ASIN"]] == ["B0ODDY0002"]
    audit = [r for r in rows("Full Audit History") if r["ASIN"] == "B0ODDY0001"]
    assert {r["Observation outcome"] for r in audit if r["Record type"] == "OBSERVATION"} == {"OK"}
    assert [r["Verification method"] for r in audit if r["Record type"] == "OBSERVATION"][
        -1
    ] == "MANUAL"


def test_manual_evidence_rows_are_append_only(env, auth):
    job = run(env, auth)
    cid = check_for(env, auth, job, "B0ODDY0001")
    submit(env, auth, job, cid, action_enabled=False)  # a mistake
    submit(env, auth, job, cid)  # corrected entry from the same source supersedes it
    with session_scope() as s:
        rows = s.scalars(select(CheckObservation).order_by(CheckObservation.id)).all()
    manual = [r for r in rows if r.source.value == "MANUAL"]
    assert len(manual) == 2 and manual[0].signals["action_enabled"] is False  # history kept
    assert env.http.get(f"/api/v1/jobs/{job}/checks/{cid}", headers=auth).json()["status"] == (
        "VERIFIED_LISTABLE"
    )
