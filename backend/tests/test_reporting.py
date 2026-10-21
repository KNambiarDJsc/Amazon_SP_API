import io
import zipfile
from datetime import datetime

import pytest
from openpyxl import load_workbook

from app.amazon.mock import NO_RESTRICTIONS, MockAmazonClient, catalog_item, demo_client
from app.db import session_scope
from app.enums import DataSource
from app.models import Job
from app.services.reporting import (
    PRODUCT_COLUMNS,
    SHEET_AUDIT,
    SHEET_LISTABLE,
    SHEET_MANUAL,
    SHEET_RESTRICTED,
    build_workbook,
)


def export(job_id):
    with session_scope() as s:
        data = build_workbook(s, s.get(Job, job_id))
    return load_workbook(io.BytesIO(data)), data


def rows(wb, sheet):
    ws = wb[sheet]
    header = [c.value for c in ws[1]]
    return [dict(zip(header, [c.value for c in r], strict=True)) for r in ws.iter_rows(min_row=2)]


def test_workbook_has_exactly_the_four_specified_sheets_with_headers(make_harness):
    h = make_harness(demo_client())
    job = h.submit(["eSSL"])
    h.drain()
    wb, _ = export(job)
    assert wb.sheetnames == [SHEET_LISTABLE, SHEET_RESTRICTED, SHEET_MANUAL, SHEET_AUDIT]
    expected = [name for name, _ in PRODUCT_COLUMNS]
    for sheet in (SHEET_LISTABLE, SHEET_RESTRICTED, SHEET_MANUAL):
        assert [c.value for c in wb[sheet][1]] == expected
    for required in (
        "Brand",
        "ASIN",
        "Product title",
        "Marketplace",
        "Condition",
        "Status",
        "Approval reason",
        "Verification method",
        "Seller account",
        "Evidence reference",
        "Checked at (UTC)",
    ):
        assert required in expected


def test_products_land_on_the_right_sheet_with_evidence(make_harness):
    h = make_harness(demo_client())
    job = h.submit(["eSSL", "Havells", "Taparia", "Oddy"])
    h.drain()
    wb, _ = export(job)

    listable = rows(wb, SHEET_LISTABLE)
    assert sorted(r["ASIN"] for r in listable) == [
        "B084FZXB48",
        "B08PDL14VR",
        "B0B63WV6CW",
        "B0BXM5CGRM",
    ]
    assert {r["Status"] for r in listable} == {"VERIFIED_LISTABLE"}
    first = next(r for r in listable if r["ASIN"] == "B08PDL14VR")
    assert first["Brand"] == "eSSL"
    assert first["Product title"] == "ESSL IDENTIX K30 PRO BIOMETRIC Machine"
    assert first["Marketplace"] == "A21TJRUUN4KGV" and first["Condition"] == "new_new"
    assert first["Verification method"] == "SP_API_RESTRICTIONS"
    assert first["Evidence reference"].startswith("db:check_observations/")
    assert isinstance(first["Checked at (UTC)"], datetime)
    assert first["Brand result (sampled ASINs only)"] == "VERIFIED_LISTABLE"

    restricted = rows(wb, SHEET_RESTRICTED)
    havells = next(r for r in restricted if r["ASIN"] == "B08Y9C5H38")
    assert havells["Status"] == "BRAND_APPROVAL_REQUIRED"
    assert havells["Approval reason"] == "You need approval to list in this brand."
    assert havells["Approval link"].startswith("https://sellercentral.amazon.in/")
    assert {r["Status"] for r in restricted if r["Brand"] == "Taparia"} == {
        "OTHER_APPROVAL_REQUIRED"
    }

    # Nothing restricted can appear as listable and vice versa.
    assert {r["ASIN"] for r in listable}.isdisjoint({r["ASIN"] for r in restricted})


def test_manual_review_lists_ambiguous_checks_and_brands_with_no_product(make_harness):
    client = MockAmazonClient(
        catalog={"Oddy": [catalog_item("B0AMBIG001", "x", "Oddy")]},  # no restriction behaviour
    )
    h = make_harness(client)
    job = h.submit(["Oddy", "Nothing"])
    h.drain()
    wb, _ = export(job)
    manual = rows(wb, SHEET_MANUAL)
    by_brand = {r["Brand"]: r for r in manual}
    assert by_brand["Nothing"]["Status"] == "NO_PRODUCT_FOUND"
    assert by_brand["Nothing"]["ASIN"] in (None, "")
    asin_row = next(r for r in manual if r["ASIN"] == "B0AMBIG001")
    assert asin_row["Status"] == "NEEDS_MANUAL_CHECK" and asin_row["Rule"] == "SPAPI_ASIN_NOT_FOUND"
    assert rows(wb, SHEET_LISTABLE) == [] and rows(wb, SHEET_RESTRICTED) == []


def test_dead_lettered_brand_appears_in_manual_review_with_the_last_error(make_harness):
    client = MockAmazonClient(catalog={"Oddy": []})
    client.catalog_behaviour = RuntimeError("boom")
    h = make_harness(client, max_brand_attempts=2)
    job = h.submit(["Oddy"])
    h.drain()
    wb, _ = export(job)
    (row,) = rows(wb, SHEET_MANUAL)
    assert row["Rule"] == "BRAND_DEAD_LETTERED" and "boom" in row["Rationale"]


def test_audit_history_has_one_row_per_observation_plus_brand_results(make_harness):
    h = make_harness(demo_client())
    job = h.submit(["Havells", "Kangaro"])
    h.drain()
    wb, _ = export(job)
    audit = rows(wb, SHEET_AUDIT)
    kinds = [r["Record type"] for r in audit]
    assert kinds.count("BRAND_RESULT") == 2 and kinds.count("OBSERVATION") == 5
    obs = [r for r in audit if r["Record type"] == "OBSERVATION"]
    assert {r["Observation outcome"] for r in obs} == {"OK"}
    assert {r["HTTP status"] for r in obs} == {200}
    brand_rows = {r["Brand"]: r for r in audit if r["Record type"] == "BRAND_RESULT"}
    assert brand_rows["Kangaro"]["Status"] == "MIXED_RESULTS"


def test_failed_attempts_are_kept_in_the_audit_history(make_harness):
    from app.amazon.errors import AmazonRequestError

    client = MockAmazonClient(
        catalog={"Oddy": [catalog_item("B0BAD00001", "x", "Oddy")]},
        restrictions={"B0BAD00001": AmazonRequestError(400, "bad")},
    )
    h = make_harness(client)
    job = h.submit(["Oddy"])
    h.drain()
    wb, _ = export(job)
    obs = [r for r in rows(wb, SHEET_AUDIT) if r["Record type"] == "OBSERVATION"]
    assert [(r["Observation outcome"], r["HTTP status"]) for r in obs] == [("HTTP_ERROR", 400)]


def test_mock_data_is_labelled_on_every_row_and_in_the_file_properties(make_harness):
    h = make_harness(demo_client())
    job = h.submit(["Havells"])
    h.drain()
    wb, _ = export(job)
    for sheet in wb.sheetnames:
        for row in rows(wb, sheet):
            assert row["Data source"].startswith("MOCK")
    assert "DEMO DATA" in wb.properties.subject


def test_live_data_is_not_labelled_mock(make_harness):
    h = make_harness(demo_client())
    job = h.submit(["Havells"])
    h.drain()
    with session_scope() as s:
        s.get(Job, job).data_source = DataSource.LIVE
    wb, _ = export(job)
    assert {r["Data source"] for r in rows(wb, SHEET_RESTRICTED)} == {"LIVE"}
    assert wb.properties.subject is None


@pytest.mark.parametrize("payload", ['=HYPERLINK("http://evil","x")', "+1+1", "-2+3", "@SUM(A1)"])
def test_text_that_looks_like_a_formula_is_written_as_text(make_harness, payload):
    brand_name = payload
    client = MockAmazonClient(
        catalog={brand_name: [catalog_item("B0EVIL0001", payload, brand_name)]},
        restrictions={"B0EVIL0001": NO_RESTRICTIONS},
    )
    h = make_harness(client)
    job = h.submit([payload])
    h.drain()
    wb, data = export(job)

    listable = rows(wb, SHEET_LISTABLE)
    assert listable[0]["Brand"] == payload and listable[0]["Product title"] == payload
    ws = wb[SHEET_LISTABLE]
    assert ws["A2"].data_type == "s" and ws["C2"].data_type == "s"  # strings, never formulas

    # And in the raw XML there is no <f> formula element anywhere.
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        for name in z.namelist():
            if name.startswith("xl/worksheets/"):
                assert b"<f>" not in z.read(name) and b"<f " not in z.read(name)


def test_export_reads_only_committed_rows_and_works_mid_job(make_harness):
    h = make_harness(demo_client(), seller_max_concurrent_brands=1)
    job = h.submit(["eSSL", "Havells"])
    h.processor.process(h.dispatcher.pop()[0])  # only eSSL has been processed
    wb, _ = export(job)
    assert len(rows(wb, SHEET_LISTABLE)) == 4 and rows(wb, SHEET_RESTRICTED) == []
    audit = rows(wb, SHEET_AUDIT)
    havells = next(r for r in audit if r["Brand"] == "Havells")
    assert havells["Status"] == "QUEUED"  # visible as not yet processed


def test_a_check_planned_but_not_yet_run_is_pending_not_a_result(make_harness, seller):
    from app.services import orchestration

    ctx = {}

    def pause_while_first_check_runs(n):
        if n == 0:
            orchestration.pause(seller.id, ctx["job"])
        return NO_RESTRICTIONS

    client = MockAmazonClient(
        catalog={"Oddy": [catalog_item(f"B0PEND000{i}", "x", "Oddy") for i in (1, 2)]},
        restrictions={"B0PEND0001": pause_while_first_check_runs, "B0PEND0002": NO_RESTRICTIONS},
    )
    h = make_harness(client)
    ctx["job"] = job = h.submit(["Oddy"])
    h.drain()

    wb, _ = export(job)
    assert [r["ASIN"] for r in rows(wb, SHEET_LISTABLE)] == ["B0PEND0001"]  # only what ran
    (pending,) = [r for r in rows(wb, SHEET_AUDIT) if r["Record type"] == "CHECK_PENDING"]
    assert pending["ASIN"] == "B0PEND0002" and pending["Status"] == "PENDING"
    assert all(r["ASIN"] != "B0PEND0002" for r in rows(wb, SHEET_MANUAL))
