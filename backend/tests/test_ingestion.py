import csv
import io
import zipfile

import pytest
from openpyxl import Workbook

from app.services.ingestion import (
    IngestionError,
    ingest,
    rows_from_csv,
    rows_from_list,
    rows_from_upload,
    rows_from_xlsx,
)

LIMITS = dict(max_brands=100, max_length=200)


def run(rows, **kw):
    return ingest(rows, **{**LIMITS, **kw})


def xlsx_bytes(grid) -> bytes:
    wb = Workbook()
    ws = wb.active
    for row in grid:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---- pasted lists ----------------------------------------------------------------------------


def test_pasted_brands_are_normalised_but_originals_preserved():
    r = run(rows_from_list(["  eSSL ", "Kangaro\t", "Taparia"]))
    assert [b.display_name for b in r.accepted] == ["eSSL", "Kangaro", "Taparia"]
    assert [b.original_name for b in r.accepted] == ["  eSSL ", "Kangaro\t", "Taparia"]
    assert [b.position for b in r.accepted] == [0, 1, 2]


def test_case_insensitive_duplicates_are_reported_with_row_numbers():
    r = run(rows_from_list(["eSSL", "Oddy", "ESSL", "  essl  ", "oddy"]))
    assert [b.display_name for b in r.accepted] == ["eSSL", "Oddy"]
    assert [(d.row, d.duplicate_of_row) for d in r.report.duplicates] == [(3, 1), (4, 1), (5, 2)]
    assert r.report.total_rows == 5 and r.report.accepted == 2


def test_blank_and_invisible_only_rows_are_reported_not_dropped_silently():
    r = run(rows_from_list(["eSSL", "", "   ", "​", "Oddy"]))
    assert r.report.empty_rows == [2, 3, 4]
    assert r.report.accepted == 2


def test_invalid_rows_are_reported_with_a_reason():
    r = run(rows_from_list(["eSSL", "-----", "x" * 300]), max_length=200)
    assert {(i.row, i.reason) for i in r.report.invalid} == {
        (2, "NO_LETTERS_OR_DIGITS"),
        (3, "TOO_LONG"),
    }
    assert len(r.accepted) == 1


def test_entries_containing_newlines_become_separate_rows():
    rows = rows_from_list(["eSSL\nKangaro\r\nTaparia", "Oddy"])
    assert [r.text for r in rows] == ["eSSL", "Kangaro", "Taparia", "Oddy"]
    assert [r.row for r in rows] == [1, 2, 3, 4]


def test_no_usable_brands_is_an_error():
    with pytest.raises(IngestionError, match="No usable"):
        run(rows_from_list(["", "  ", "---"]))


def test_job_size_limit_is_enforced_on_accepted_brands_not_raw_rows():
    ok = run(rows_from_list(["a", "A", "", "b"]), max_brands=2)  # duplicates/blanks do not count
    assert ok.report.accepted == 2
    with pytest.raises(IngestionError, match="Too many brands"):
        run(rows_from_list(["a", "b", "c"]), max_brands=2)


def test_unicode_brands_survive():
    r = run(rows_from_list(["Café Noir", "ＥＳＳＬ", "पतंजलि"]))
    assert [b.display_name for b in r.accepted] == ["Café Noir", "ESSL", "पतंजलि"]


# ---- CSV -------------------------------------------------------------------------------------


def test_csv_with_brand_header_uses_that_column_and_spreadsheet_row_numbers():
    data = b"Status,Brand Name,Notes\nPending,eSSL,x\nPending,Kangaro,y\n"
    rows = rows_from_csv(data)
    assert [(r.row, r.text) for r in rows] == [(2, "eSSL"), (3, "Kangaro")]


def test_csv_without_header_uses_first_column_for_every_row():
    rows = rows_from_csv(b"eSSL\nKangaro\n")
    assert [(r.row, r.text) for r in rows] == [(1, "eSSL"), (2, "Kangaro")]


def test_csv_quoted_commas_and_bom_and_semicolons():
    buf = io.StringIO()
    csv.writer(buf).writerow(["Johnson, Inc", "x"])
    assert rows_from_csv(buf.getvalue().encode())[0].text == "Johnson, Inc"
    assert rows_from_csv("﻿Brand\neSSL\n".encode())[0].text == "eSSL"
    assert [r.text for r in rows_from_csv(b"Brand;Status\neSSL;Pending\nOddy;Pending\n")] == [
        "eSSL",
        "Oddy",
    ]


def test_csv_falls_back_to_cp1252_for_excel_on_windows():
    data = "Brand\nCafé".encode("cp1252")
    assert rows_from_csv(data)[0].text == "Café"


# ---- XLSX ------------------------------------------------------------------------------------


def test_xlsx_matches_the_client_example_layout():
    data = xlsx_bytes(
        [
            ["Brand Name", "Status", "Product Name", "ASIN", "Notes"],
            ["eSSL", "Pending", None, None, None],
            ["Kangaro", "Pending", None, None, None],
            ["Taparia", "Pending", None, None, None],
        ]
    )
    rows = rows_from_xlsx(data)
    assert [(r.row, r.text) for r in rows] == [(2, "eSSL"), (3, "Kangaro"), (4, "Taparia")]


def test_xlsx_numeric_cells_and_blank_cells():
    rows = rows_from_xlsx(xlsx_bytes([["Brand"], [3.0], [None], ["3M"]]))
    assert [r.text for r in rows] == ["3", "", "3M"]


def test_xlsx_finds_brand_column_anywhere_in_the_header():
    rows = rows_from_xlsx(xlsx_bytes([["SKU", "brand"], ["1", "eSSL"], ["2", "Oddy"]]))
    assert [r.text for r in rows] == ["eSSL", "Oddy"]


def test_xlsx_formulas_are_never_evaluated():
    # A formula cell with no cached value reads as empty and is reported as an empty row.
    r = run(rows_from_xlsx(xlsx_bytes([["Brand"], ["=1+1"], ["eSSL"]])))
    assert r.report.empty_rows == [2] and [b.display_name for b in r.accepted] == ["eSSL"]


def test_xlsx_text_cell_starting_with_equals_is_kept_as_text():
    wb = Workbook()
    ws = wb.active
    ws.append(["Brand"])
    ws.append(["placeholder"])
    ws["A2"].value = "=cmd|' /C calc'!A0"
    ws["A2"].data_type = "s"  # a genuine text cell that happens to start with "="
    buf = io.BytesIO()
    wb.save(buf)
    assert rows_from_xlsx(buf.getvalue())[0].text.startswith("=cmd")


def test_corrupt_and_oversized_uploads_are_rejected_cleanly():
    with pytest.raises(IngestionError, match="not a valid .xlsx"):
        rows_from_xlsx(b"PK\x03\x04 definitely not a zip")
    with pytest.raises(IngestionError, match="larger than"):
        rows_from_upload("b.csv", b"x" * 2048, max_bytes=1024)
    with pytest.raises(IngestionError, match="empty"):
        rows_from_upload("b.csv", b"  \n", max_bytes=1024)
    with pytest.raises(IngestionError, match="xls"):
        rows_from_upload("b.xls", b"\xd0\xcf\x11\xe0 old format", max_bytes=1024)


def test_zip_bomb_is_refused_before_parsing():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("xl/big.xml", b"0" * (60 * 1024 * 1024))
    assert len(buf.getvalue()) < 1024 * 1024  # tiny on the wire, huge when expanded
    with pytest.raises(IngestionError, match="too large once decompressed"):
        rows_from_upload("bomb.xlsx", buf.getvalue(), max_bytes=5 * 1024 * 1024)


def test_upload_dispatches_by_content_not_just_extension():
    data = xlsx_bytes([["Brand"], ["eSSL"]])
    assert rows_from_upload("misnamed.csv", data, max_bytes=10**7)[0].text == "eSSL"
