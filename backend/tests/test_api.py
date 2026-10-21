import io
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from support import make_runtime

from app.amazon import errors as err
from app.amazon.mock import demo_client
from app.api.main import create_app
from app.auth import create_seller_account
from app.config import get_settings
from app.db import session_scope
from app.services.tick import run_until_idle

DEMO = ["eSSL", "Havells", "Kangaro", "Taparia", "Oddy"]
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
def client(engine):
    return TestClient(create_app())


@pytest.fixture
def key(engine) -> str:
    with session_scope() as s:
        _, k = create_seller_account(s, name="Acme Traders", amazon_seller_id="A1ACME")
    return k


@pytest.fixture
def other_key(engine) -> str:
    with session_scope() as s:
        _, k = create_seller_account(s, name="Rival Co")
    return k


def auth(k: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {k}"}


def xlsx_file(rows) -> bytes:
    wb = Workbook()
    for row in rows:
        wb.active.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def create(client, key, brands=DEMO, **extra):
    response = client.post("/api/v1/jobs", json={"brands": brands, **extra}, headers=auth(key))
    assert response.status_code == 201, response.text
    return response.json()


def run_all(clock, client_obj=None):
    run_until_idle(make_runtime(client_obj or demo_client(), clock))


# ---- authentication ------------------------------------------------------------------------


def test_health_is_open_and_reports_mode(client):
    response = client.get("/health")
    assert response.status_code == 200
    assert response.json() == {"status": "ok", "amazon_mode": "mock"}


@pytest.mark.parametrize(
    "headers", [{}, {"Authorization": "Bearer nope"}, {"Authorization": "Basic abc"}]
)
def test_every_job_endpoint_rejects_missing_or_wrong_keys(client, key, headers):
    job = create(client, key)
    for method, path in [
        ("get", "/api/v1/jobs"),
        ("post", "/api/v1/jobs"),
        ("get", f"/api/v1/jobs/{job['id']}"),
        ("post", f"/api/v1/jobs/{job['id']}/pause"),
        ("get", f"/api/v1/jobs/{job['id']}/export.xlsx"),
    ]:
        response = client.request(method, path, headers=headers, json={"brands": ["x"]})
        assert response.status_code == 401, (method, path)
        assert response.headers["www-authenticate"] == "Bearer"


def test_a_deactivated_seller_is_locked_out(client, key):
    with session_scope() as s:
        from app.auth import authenticate

        authenticate(s, key).is_active = False
    assert client.get("/api/v1/jobs", headers=auth(key)).status_code == 401


# ---- creating jobs -------------------------------------------------------------------------


def test_create_job_returns_the_ingestion_report_and_demo_notice(client, key):
    job = create(client, key, ["eSSL", "essl", "", "Kangaro"], name="Q4 sourcing")
    assert job["status"] == "QUEUED"
    assert job["name"] == "Q4 sourcing"
    assert job["data_source"] == "MOCK"
    assert "DEMO DATA" in job["notice"]
    assert job["ingestion_report"]["accepted"] == 2
    assert [d["row"] for d in job["ingestion_report"]["duplicates"]] == [2]
    assert job["ingestion_report"]["empty_rows"] == [3]
    assert job["progress"]["brands_total"] == 2


@pytest.mark.parametrize(
    "body",
    [
        {"brands": []},
        {"brands": ["   ", "---"]},
        {"brands": ["a"], "settings": {"max_asins_per_brand": 11}},
        {"brands": ["a"], "settings": {"conditions": ["new_new", "new_new"]}},
        {"brands": ["a"], "settings": {"conditions": ["bogus"]}},
        {"brands": ["a"], "settings": {"min_sample_size": 6, "max_asins_per_brand": 5}},
        {"name": "x" * 201, "brands": ["a"]},
    ],
)
def test_invalid_job_requests_are_rejected_with_422(client, key, body):
    assert client.post("/api/v1/jobs", json=body, headers=auth(key)).status_code == 422


def test_too_many_brands_is_rejected_not_truncated(client, key, monkeypatch):
    monkeypatch.setenv("MAX_BRANDS_PER_JOB", "3")
    get_settings.cache_clear()
    try:
        response = client.post("/api/v1/jobs", json={"brands": list("abcd")}, headers=auth(key))
    finally:
        get_settings.cache_clear()
    assert response.status_code == 422 and "exceeds the limit" in response.json()["detail"]


def test_unconfigured_live_mode_refuses_jobs_instead_of_using_fixtures(client, key, monkeypatch):
    monkeypatch.setenv("AMAZON_MODE", "live")
    for name in ("LWA_CLIENT_ID", "LWA_CLIENT_SECRET", "CREDENTIALS_ENCRYPTION_KEYS"):
        monkeypatch.delenv(name, raising=False)
    get_settings.cache_clear()
    try:
        response = client.post("/api/v1/jobs", json={"brands": ["eSSL"]}, headers=auth(key))
    finally:
        get_settings.cache_clear()
    assert response.status_code == 503
    assert "LWA_CLIENT_ID is not set" in response.json()["detail"]


def test_preview_reports_what_would_be_accepted_without_creating_a_job(client, key):
    response = client.post(
        "/api/v1/jobs/preview", json={"brands": ["eSSL", "ESSL", "Oddy"]}, headers=auth(key)
    )
    assert response.status_code == 200
    assert response.json()["brands"] == ["eSSL", "Oddy"]
    assert len(response.json()["report"]["duplicates"]) == 1
    assert client.get("/api/v1/jobs", headers=auth(key)).json()["total"] == 0


# ---- uploads -------------------------------------------------------------------------------


def test_upload_the_clients_example_workbook(client, key):
    data = xlsx_file(
        [
            ["Brand Name", "Status", "Product Name", "ASIN", "Notes"],
            ["eSSL", "Pending"],
            ["Kangaro", "Pending"],
            ["Taparia", "Pending"],
            ["Oddy", "Pending"],
            ["Plantex", "Pending"],
        ]
    )
    response = client.post(
        "/api/v1/jobs/upload",
        files={"file": ("brands.xlsx", data, XLSX)},
        data={"name": "From workbook"},
        headers=auth(key),
    )
    assert response.status_code == 201, response.text
    job = response.json()
    assert job["source_filename"] == "brands.xlsx"
    assert job["progress"]["brands_total"] == 5 and job["name"] == "From workbook"


def test_upload_csv_with_job_settings(client, key):
    response = client.post(
        "/api/v1/jobs/upload",
        files={"file": ("b.csv", b"Brand\neSSL\nHavells\n", "text/csv")},
        data={"settings": json.dumps({"max_asins_per_brand": 2, "min_sample_size": 2})},
        headers=auth(key),
    )
    assert response.status_code == 201
    assert response.json()["settings"]["max_asins_per_brand"] == 2


@pytest.mark.parametrize(
    "filename,content,expected",
    [
        ("evil.exe", b"MZ", "unsupported"),
        ("fake.xlsx", b"not a zip at all", "not a valid .xlsx"),
        ("empty.csv", b"Brand\n", "no valid brand names"),
    ],
)
def test_bad_uploads_are_rejected_clearly(client, key, filename, content, expected):
    response = client.post(
        "/api/v1/jobs/upload", files={"file": (filename, content)}, headers=auth(key)
    )
    assert response.status_code == 422
    assert expected in response.json()["detail"]


def test_oversized_upload_is_rejected_before_parsing(client, key, monkeypatch):
    monkeypatch.setenv("MAX_UPLOAD_BYTES", "100")
    get_settings.cache_clear()
    try:
        response = client.post(
            "/api/v1/jobs/upload",
            files={"file": ("b.csv", b"Brand\n" + b"x\n" * 500)},
            headers=auth(key),
        )
    finally:
        get_settings.cache_clear()
    assert response.status_code == 413


def test_invalid_settings_json_on_upload_is_422(client, key):
    response = client.post(
        "/api/v1/jobs/upload",
        files={"file": ("b.csv", b"Brand\neSSL\n")},
        data={"settings": '{"max_asins_per_brand": 99}'},
        headers=auth(key),
    )
    assert response.status_code == 422


# ---- reading results -----------------------------------------------------------------------


def test_full_flow_from_submission_to_checks_and_evidence(client, key, clock):
    job = create(client, key)
    run_all(clock)

    done = client.get(f"/api/v1/jobs/{job['id']}", headers=auth(key)).json()
    assert done["status"] == "COMPLETED"
    assert done["progress"]["percent_complete"] == 100.0
    assert done["progress"]["verified_listable"] == 6 and done["progress"]["restricted"] == 6

    brands = client.get(f"/api/v1/jobs/{job['id']}/brands", headers=auth(key)).json()
    assert {b["display_name"]: b["brand_status"] for b in brands["items"]} == {
        "eSSL": "VERIFIED_LISTABLE",
        "Havells": "BRAND_APPROVAL_REQUIRED",
        "Kangaro": "MIXED_RESULTS",
        "Taparia": "OTHER_APPROVAL_REQUIRED",
        "Oddy": "NO_PRODUCT_FOUND",
    }
    assert brands["items"][0]["brand_summary"]["brand_wide_claim"] is False

    gated = client.get(
        f"/api/v1/jobs/{job['id']}/checks",
        params={"status": "BRAND_APPROVAL_REQUIRED"},
        headers=auth(key),
    ).json()
    assert gated["total"] == 3
    havells = next(c for c in gated["items"] if c["asin"] == "B08Y9C5H38")
    assert havells["messages"] == ["You need approval to list in this brand."]
    assert havells["approval_links"][0]["title"] == "Apply to sell"

    detail = client.get(
        f"/api/v1/jobs/{job['id']}/checks/{havells['id']}", headers=auth(key)
    ).json()
    assert detail["candidate"]["brand_attribute"] == "Havells"
    assert detail["rules_version"]
    (obs,) = detail["observations"]
    assert obs["source"] == "SP_API_RESTRICTIONS" and obs["outcome"] == "OK"
    assert obs["raw"]["restrictions"][0]["reasons"][0]["reasonCode"] == "APPROVAL_REQUIRED"


def test_check_filters_search_and_pagination(client, key, clock):
    job = create(client, key)
    run_all(clock)
    base = f"/api/v1/jobs/{job['id']}/checks"

    everything = client.get(base, headers=auth(key)).json()
    assert everything["total"] == 12
    page = client.get(base, params={"limit": 5, "offset": 10}, headers=auth(key)).json()
    assert len(page["items"]) == 2 and page["total"] == 12 and page["offset"] == 10

    by_asin = client.get(base, params={"q": "b08pdl"}, headers=auth(key)).json()
    assert [c["asin"] for c in by_asin["items"]] == ["B08PDL14VR"]
    by_title = client.get(base, params={"q": "stapler"}, headers=auth(key)).json()
    assert [c["asin"] for c in by_title["items"]] == ["B0MOCKKG01"]
    by_brand = client.get(base, params={"q": "taparia"}, headers=auth(key)).json()
    assert by_brand["total"] == 3

    # A user's % and _ are literals, not wildcards.
    assert client.get(base, params={"q": "%"}, headers=auth(key)).json()["total"] == 0
    assert client.get(base, params={"q": "B0_"}, headers=auth(key)).json()["total"] == 0

    one_brand = client.get(
        f"/api/v1/jobs/{job['id']}/brands", params={"q": "kang"}, headers=auth(key)
    ).json()
    brand_id = one_brand["items"][0]["id"]
    scoped = client.get(base, params={"brand_id": brand_id}, headers=auth(key)).json()
    assert scoped["total"] == 3

    bad = client.get(base, params={"status": "NOT_A_STATUS"}, headers=auth(key))
    assert bad.status_code == 422
    assert client.get(base, params={"limit": 0}, headers=auth(key)).status_code == 422
    assert client.get(base, params={"limit": 1000}, headers=auth(key)).status_code == 422


def test_brand_filters_by_result_and_run_status(client, key, clock):
    job = create(client, key)
    run_all(clock)
    base = f"/api/v1/jobs/{job['id']}/brands"
    listable = client.get(base, params={"result": "VERIFIED_LISTABLE"}, headers=auth(key)).json()
    assert [b["display_name"] for b in listable["items"]] == ["eSSL"]
    completed = client.get(base, params={"status": "COMPLETED"}, headers=auth(key)).json()
    assert completed["total"] == 5


def test_audit_trail_is_readable(client, key, clock):
    job = create(client, key)
    run_all(clock)
    events = client.get(f"/api/v1/jobs/{job['id']}/audit", headers=auth(key)).json()
    types = [e["event_type"] for e in events]
    assert types[0] == "job_created" and types[-1] == "job_completed"
    assert types.count("brand_completed") == 5


def test_job_list_is_paginated_newest_first(client, key):
    ids = [create(client, key, ["eSSL"], name=f"job {n}")["id"] for n in range(3)]
    listing = client.get("/api/v1/jobs", params={"limit": 2}, headers=auth(key)).json()
    assert listing["total"] == 3 and len(listing["items"]) == 2
    assert listing["items"][0]["id"] == ids[-1]


# ---- the export ----------------------------------------------------------------------------


def test_export_downloads_a_valid_workbook(client, key, clock):
    job = create(client, key, name="Q4/Sourcing: ../x")
    run_all(clock)

    response = client.get(f"/api/v1/jobs/{job['id']}/export.xlsx", headers=auth(key))

    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX
    assert response.headers["cache-control"] == "no-store"
    disposition = response.headers["content-disposition"]
    assert disposition.startswith('attachment; filename="') and "/" not in disposition
    wb = load_workbook(io.BytesIO(response.content))
    assert wb.sheetnames[:4] == [
        "Verified Listable Products",
        "Restricted Products",
        "Manual Review",
        "Full Audit History",
    ]
    assert wb["Verified Listable Products"]["A1"].value.startswith("DEMO DATA")
    assert wb["Verified Listable Products"].max_row == 2 + 6


# ---- controls ------------------------------------------------------------------------------


def test_pause_resume_cancel_over_http(client, key, clock):
    job = create(client, key)
    base = f"/api/v1/jobs/{job['id']}"

    paused = client.post(f"{base}/pause", headers=auth(key)).json()
    assert paused["status"] == "PAUSED" and paused["pause_reason"] == "Paused by user"
    run_all(clock)
    assert client.get(base, headers=auth(key)).json()["progress"]["brands_completed"] == 0

    assert client.post(f"{base}/resume", headers=auth(key)).json()["status"] == "QUEUED"
    assert client.post(f"{base}/resume", headers=auth(key)).status_code == 200  # idempotent

    cancelled = client.post(f"{base}/cancel", headers=auth(key)).json()
    assert cancelled["status"] == "CANCELLED"
    assert cancelled["progress"]["brands_cancelled"] == 5
    assert client.post(f"{base}/resume", headers=auth(key)).status_code == 409
    assert client.post(f"{base}/pause", headers=auth(key)).status_code == 409


def test_retry_failed_over_http(client, key, clock):
    job = create(client, key, ["Kangaro"])
    flaky = demo_client()
    flaky.restrictions["B0MOCKKG02"] = err.AmazonTransientError("down")
    run_all(clock, flaky)
    failed = client.get(f"/api/v1/jobs/{job['id']}", headers=auth(key)).json()
    assert failed["progress"]["brands_failed"] == 1

    flaky.restrictions["B0MOCKKG02"] = {"restrictions": []}
    retried = client.post(f"/api/v1/jobs/{job['id']}/retry-failed", headers=auth(key)).json()
    assert retried["progress"]["brands_failed"] == 0 and retried["status"] == "RUNNING"
    run_all(clock, flaky)
    final = client.get(f"/api/v1/jobs/{job['id']}", headers=auth(key)).json()
    assert final["status"] == "COMPLETED" and final["progress"]["needs_manual_check"] == 0


# ---- tenant isolation ----------------------------------------------------------------------


def test_another_sellers_job_does_not_exist_as_far_as_you_can_tell(client, key, other_key, clock):
    job = create(client, key)
    run_all(clock)
    check_id = client.get(f"/api/v1/jobs/{job['id']}/checks", headers=auth(key)).json()["items"][0][
        "id"
    ]
    base = f"/api/v1/jobs/{job['id']}"

    for method, path in [
        ("get", base),
        ("post", f"{base}/pause"),
        ("post", f"{base}/resume"),
        ("post", f"{base}/cancel"),
        ("post", f"{base}/retry-failed"),
        ("get", f"{base}/brands"),
        ("get", f"{base}/checks"),
        ("get", f"{base}/checks/{check_id}"),
        ("get", f"{base}/audit"),
        ("get", f"{base}/export.xlsx"),
    ]:
        response = client.request(method, path, headers=auth(other_key))
        assert response.status_code == 404, (method, path, response.status_code)
        assert response.json() == {
            "detail": "job not found"
        }  # same as for a job that never existed

    assert client.get("/api/v1/jobs", headers=auth(other_key)).json() == {
        "items": [],
        "total": 0,
        "limit": 50,
        "offset": 0,
    }
    # ...and the owner's job was not touched by any of the attempts above.
    assert client.get(base, headers=auth(key)).json()["status"] == "COMPLETED"


def test_a_check_id_from_another_job_is_not_found_even_for_the_owner(client, key, clock):
    first = create(client, key, ["Havells"])
    second = create(client, key, ["eSSL"])
    run_all(clock)
    check_of_first = client.get(f"/api/v1/jobs/{first['id']}/checks", headers=auth(key)).json()[
        "items"
    ][0]["id"]
    response = client.get(f"/api/v1/jobs/{second['id']}/checks/{check_of_first}", headers=auth(key))
    assert response.status_code == 404


def test_malformed_ids_are_422_not_500(client, key):
    assert client.get("/api/v1/jobs/not-a-uuid", headers=auth(key)).status_code == 422
