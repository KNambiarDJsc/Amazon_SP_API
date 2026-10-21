import io
import os
import uuid
from datetime import timedelta

import pytest
from apienv import Env
from conftest import IS_POSTGRES
from openpyxl import load_workbook
from sqlalchemy import select

from app.db import session_scope
from app.domain.observations import utcnow
from app.enums import ExportStatus
from app.models import AuditEvent, Export
from app.services import exports

SHEETS = [
    "Verified Listable Products",
    "Restricted Products",
    "Manual Review",
    "Full Audit History",
]


@pytest.fixture
def env(engine, tmp_path):
    e = Env()
    e.cfg.export_dir = str(tmp_path / "exports")
    e.cfg.export_sync_max_checks = 3  # force the asynchronous path for the demo job (9 checks)
    return e


@pytest.fixture
def auth(env):
    return env.new_seller()[1]


def demo_job(env, auth):
    job = env.http.post(
        "/api/v1/jobs", json={"brands": ["eSSL", "Havells", "Kangaro"]}, headers=auth
    ).json()["id"]
    env.drain()
    return job


def run_worker(env):
    """What the export worker does for each queued message."""
    while env.dispatcher.export_messages:
        exports.run_export(env.dispatcher.export_messages.pop(0), env.cfg)


# ---- the size threshold ----------------------------------------------------------------------


def test_small_jobs_export_inline_and_large_ones_are_pointed_at_the_async_route(env, auth):
    job = demo_job(env, auth)
    r = env.http.get(f"/api/v1/jobs/{job}/export.xlsx", headers=auth)
    assert r.status_code == 409
    assert "9 checks" in r.json()["detail"] and f"/api/v1/jobs/{job}/exports" in r.json()["detail"]
    env.cfg.export_sync_max_checks = 100
    ok = env.http.get(f"/api/v1/jobs/{job}/export.xlsx", headers=auth)
    assert ok.status_code == 200 and load_workbook(io.BytesIO(ok.content)).sheetnames == SHEETS


# ---- the asynchronous flow -------------------------------------------------------------------


def test_request_build_poll_download(env, auth):
    job = demo_job(env, auth)
    r = env.http.post(f"/api/v1/jobs/{job}/exports", headers=auth)
    assert r.status_code == 202
    body = r.json()
    assert body["status"] == "QUEUED" and body["download_url"] is None
    assert env.dispatcher.export_messages == [uuid.UUID(body["id"])]
    url = f"/api/v1/jobs/{job}/exports/{body['id']}"
    assert env.http.get(url + "/download", headers=auth).status_code == 409  # not ready yet

    run_worker(env)
    ready = env.http.get(url, headers=auth).json()
    assert ready["status"] == "READY" and ready["size_bytes"] > 1000 and ready["checks_count"] == 9
    assert ready["expires_at"] and ready["download_url"] == url + "/download"
    assert ready["error"] is None

    dl = env.http.get(ready["download_url"], headers=auth)
    assert dl.status_code == 200
    assert dl.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert f"eligibility-{job[:8]}.xlsx" in dl.headers["content-disposition"]
    assert dl.headers["cache-control"] == "no-store"
    wb = load_workbook(io.BytesIO(dl.content))
    assert wb.sheetnames == SHEETS and wb["Verified Listable Products"].max_row == 7
    types = [
        e["event_type"] for e in env.http.get(f"/api/v1/jobs/{job}/events", headers=auth).json()
    ]
    assert "EXPORT_REQUESTED" in types and "EXPORT_READY" in types


def test_the_file_is_named_by_id_never_by_user_input_and_written_atomically(env, auth, tmp_path):
    job = demo_job(env, auth)
    eid = env.http.post(f"/api/v1/jobs/{job}/exports", headers=auth).json()["id"]
    run_worker(env)
    files = sorted(os.listdir(tmp_path / "exports"))
    assert files == [f"{eid}.xlsx"]  # no leftover .part file


def test_a_second_request_while_one_is_in_flight_returns_the_same_export(env, auth):
    job = demo_job(env, auth)
    a = env.http.post(f"/api/v1/jobs/{job}/exports", headers=auth).json()["id"]
    b = env.http.post(f"/api/v1/jobs/{job}/exports", headers=auth).json()["id"]
    assert a == b and len(env.dispatcher.export_messages) == 1
    run_worker(env)
    c = env.http.post(f"/api/v1/jobs/{job}/exports", headers=auth).json()["id"]
    assert c != a  # once finished, a new request builds a fresh report
    assert len(env.http.get(f"/api/v1/jobs/{job}/exports", headers=auth).json()) == 2


def test_running_the_same_message_twice_builds_once(env, auth):
    job = demo_job(env, auth)
    eid = uuid.UUID(env.http.post(f"/api/v1/jobs/{job}/exports", headers=auth).json()["id"])
    exports.run_export(eid, env.cfg)
    with session_scope() as s:
        first = s.get(Export, eid).finished_at
    exports.run_export(eid, env.cfg)  # redelivery
    with session_scope() as s:
        assert s.get(Export, eid).finished_at == first


def test_a_failed_build_is_recorded_and_can_be_requested_again(env, auth, monkeypatch):
    job = demo_job(env, auth)
    eid = env.http.post(f"/api/v1/jobs/{job}/exports", headers=auth).json()["id"]

    def boom(*_a, **_k):
        raise RuntimeError("disk full")

    monkeypatch.setattr(exports, "build_workbook", boom)
    run_worker(env)
    failed = env.http.get(f"/api/v1/jobs/{job}/exports/{eid}", headers=auth).json()
    assert failed["status"] == "FAILED" and "RuntimeError: disk full" in failed["error"]
    assert failed["download_url"] is None
    monkeypatch.undo()
    again = env.http.post(f"/api/v1/jobs/{job}/exports", headers=auth).json()
    assert again["id"] != eid and again["status"] == "QUEUED"
    run_worker(env)
    assert (
        env.http.get(f"/api/v1/jobs/{job}/exports/{again['id']}", headers=auth).json()["status"]
        == "READY"
    )


def test_a_queue_that_is_down_fails_the_request_cleanly(env, auth):
    job = demo_job(env, auth)

    def down(_id):
        raise ConnectionError("redis down")

    env.dispatcher.enqueue_export = down
    with pytest.raises(ConnectionError):
        env.http.post(f"/api/v1/jobs/{job}/exports", headers=auth)
    with session_scope() as s:
        (export,) = s.scalars(select(Export)).all()
        assert export.status == ExportStatus.FAILED and "queue unavailable" in export.error


# ---- expiry and housekeeping -----------------------------------------------------------------


def test_expired_exports_cannot_be_downloaded_and_are_purged_with_their_files(env, auth, tmp_path):
    job = demo_job(env, auth)
    eid = env.http.post(f"/api/v1/jobs/{job}/exports", headers=auth).json()["id"]
    run_worker(env)
    path = tmp_path / "exports" / f"{eid}.xlsx"
    assert path.exists()
    with session_scope() as s:
        s.get(Export, uuid.UUID(eid)).expires_at = utcnow() - timedelta(minutes=1)
    gone = env.http.get(f"/api/v1/jobs/{job}/exports/{eid}/download", headers=auth)
    assert gone.status_code == 410 and "expired" in gone.json()["detail"]

    assert exports.purge(env.cfg) == {"purged": 1, "failed_stale": 0}
    assert not path.exists()
    assert env.http.get(f"/api/v1/jobs/{job}/exports/{eid}", headers=auth).status_code == 404


def test_a_deleted_file_is_reported_as_expired_not_a_server_error(env, auth, tmp_path):
    job = demo_job(env, auth)
    eid = env.http.post(f"/api/v1/jobs/{job}/exports", headers=auth).json()["id"]
    run_worker(env)
    (tmp_path / "exports" / f"{eid}.xlsx").unlink()
    assert (
        env.http.get(f"/api/v1/jobs/{job}/exports/{eid}/download", headers=auth).status_code == 410
    )


def test_an_export_whose_worker_died_is_failed_after_the_stale_window(env, auth):
    job = demo_job(env, auth)
    eid = uuid.UUID(env.http.post(f"/api/v1/jobs/{job}/exports", headers=auth).json()["id"])
    with session_scope() as s:  # claimed, then the worker vanished
        row = s.get(Export, eid)
        row.status, row.started_at = ExportStatus.RUNNING, utcnow() - timedelta(minutes=31)
    assert exports.purge(env.cfg)["failed_stale"] == 1
    with session_scope() as s:
        assert s.get(Export, eid).status == ExportStatus.FAILED
        assert "stopped" in s.get(Export, eid).error
    # and the job can request a new one
    assert env.http.post(f"/api/v1/jobs/{job}/exports", headers=auth).status_code == 202


def test_old_failed_exports_are_purged(env, auth):
    job = demo_job(env, auth)
    eid = uuid.UUID(env.http.post(f"/api/v1/jobs/{job}/exports", headers=auth).json()["id"])
    with session_scope() as s:
        row = s.get(Export, eid)
        row.status, row.finished_at = ExportStatus.FAILED, utcnow() - timedelta(hours=25)
    assert exports.purge(env.cfg)["purged"] == 1


def test_recovery_runs_the_purge(env, auth):
    from app.services.dispatch import recover

    job = demo_job(env, auth)
    eid = uuid.UUID(env.http.post(f"/api/v1/jobs/{job}/exports", headers=auth).json()["id"])
    run_worker(env)
    with session_scope() as s:
        s.get(Export, eid).expires_at = utcnow() - timedelta(seconds=1)
    assert recover(env.dispatcher, env.cfg).exports_purged == 1


# ---- isolation -------------------------------------------------------------------------------


def test_exports_belong_to_one_seller(env, auth):
    job = demo_job(env, auth)
    eid = env.http.post(f"/api/v1/jobs/{job}/exports", headers=auth).json()["id"]
    run_worker(env)
    _, other = env.new_seller("Other")
    for method, url in [
        ("post", f"/api/v1/jobs/{job}/exports"),
        ("get", f"/api/v1/jobs/{job}/exports"),
        ("get", f"/api/v1/jobs/{job}/exports/{eid}"),
        ("get", f"/api/v1/jobs/{job}/exports/{eid}/download"),
    ]:
        assert getattr(env.http, method)(url, headers=other).status_code == 404, url
    # an export id from another job of the same seller is not found under this job
    second = env.http.post("/api/v1/jobs", json={"brands": ["Oddy"]}, headers=auth).json()["id"]
    assert env.http.get(f"/api/v1/jobs/{second}/exports/{eid}", headers=auth).status_code == 404
    assert (
        env.http.get(f"/api/v1/jobs/{job}/exports/{uuid.uuid4()}", headers=auth).status_code == 404
    )


# ---- consistency -----------------------------------------------------------------------------


@pytest.mark.skipif(not IS_POSTGRES, reason="isolation levels are a PostgreSQL feature")
def test_the_report_is_built_from_a_single_snapshot(env, auth, monkeypatch):
    from sqlalchemy import text

    job = demo_job(env, auth)
    eid = uuid.UUID(env.http.post(f"/api/v1/jobs/{job}/exports", headers=auth).json()["id"])
    seen = []
    real = exports.build_workbook

    def spy(session, job_row):
        seen.append(session.execute(text("show transaction_isolation")).scalar())
        return real(session, job_row)

    monkeypatch.setattr(exports, "build_workbook", spy)
    exports.run_export(eid, env.cfg)
    assert seen == ["repeatable read"]


def test_the_dramatiq_runtime_registers_a_separate_export_queue():
    from dramatiq.brokers.stub import StubBroker

    from app.config import Settings
    from app.services.rate_limit import InMemoryTokenBucket
    from app.workers.runtime import EXPORT_QUEUE, QUEUE, build_runtime

    runtime = build_runtime(StubBroker(), Settings(database_url="x"), limiter=InMemoryTokenBucket())
    assert runtime.dispatcher._export_actor.queue_name == EXPORT_QUEUE != QUEUE
    assert runtime.dispatcher._export_actor.actor_name == "build_export"


def test_audit_events_for_exports_stay_inside_the_job(env, auth):
    job = demo_job(env, auth)
    env.http.post(f"/api/v1/jobs/{job}/exports", headers=auth)
    with session_scope() as s:
        events = s.scalars(
            select(AuditEvent).where(AuditEvent.event_type == "EXPORT_REQUESTED")
        ).all()
        assert [str(e.job_id) for e in events] == [job]
