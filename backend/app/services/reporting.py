"""Excel export, generated on demand from committed database rows.

Four sheets: Verified Listable Products, Restricted Products, Manual Review, Full Audit History.
Cell text can originate from brand names a client typed and product titles Amazon returned, so
any string that a spreadsheet could interpret as a formula is written as plain text.
"""

import io
from collections.abc import Iterable, Sequence
from datetime import UTC, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.enums import (
    BrandStatus,
    CheckState,
    DataSource,
    EligibilityStatus,
)
from app.models import (
    CatalogCandidate,
    CheckObservation,
    EligibilityCheck,
    Job,
    JobBrand,
    SellerAccount,
)

SHEET_LISTABLE = "Verified Listable Products"
SHEET_RESTRICTED = "Restricted Products"
SHEET_MANUAL = "Manual Review"
SHEET_AUDIT = "Full Audit History"

PRODUCT_COLUMNS = [
    ("Brand", 22),
    ("ASIN", 14),
    ("Product title", 60),
    ("Marketplace", 16),
    ("Condition", 14),
    ("Status", 26),
    ("Approval reason", 48),
    ("Approval link", 40),
    ("Verification method", 22),
    ("Seller account", 18),
    ("Evidence reference", 30),
    ("Checked at (UTC)", 20),
    ("Brand result (sampled ASINs only)", 30),
    ("Rule", 28),
    ("Rationale", 60),
    ("Data source", 18),
]
AUDIT_EXTRA = [("Record type", 16), ("Observation outcome", 20), ("HTTP status", 12)]

_FORMULA_FIRST = frozenset("=+-@\t\r")
RESTRICTED = (
    EligibilityStatus.BRAND_APPROVAL_REQUIRED,
    EligibilityStatus.OTHER_APPROVAL_REQUIRED,
    EligibilityStatus.NOT_ELIGIBLE,
)
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")


def _excel_datetime(value: datetime | None) -> datetime | None:
    # Excel has no time zones; the column header says UTC.
    return value.astimezone(UTC).replace(tzinfo=None) if value else None


def _cell(ws, value: Any):
    """A value safe to append in write-only mode. Plain values are appended as they are (this is
    the hot path: millions of cells); only text a spreadsheet could read as a formula is wrapped
    so it is stored as text."""
    if isinstance(value, datetime):
        return _excel_datetime(value)  # openpyxl gives datetimes a date format on its own
    if isinstance(value, str) and value[:1] in _FORMULA_FIRST:
        cell = WriteOnlyCell(ws, value=value)
        cell.data_type = "s"  # force text; never a formula
        return cell
    return value


def _header(ws, columns: Sequence[tuple[str, int]]) -> None:
    for index, (_, width) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"
    cells = []
    for name, _ in columns:
        cell = WriteOnlyCell(ws, value=name)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cells.append(cell)
    ws.append(cells)


def _finish(ws, columns: Sequence[tuple[str, int]], rows: int) -> None:
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(rows + 1, 1)}"


def _write(ws, columns, rows: Iterable[Sequence[Any]]) -> int:
    _header(ws, columns)
    count = 0
    for row in rows:
        ws.append([_cell(ws, v) for v in row])
        count += 1
    _finish(ws, columns, count)
    return count


def _reason_text(check: EligibilityCheck) -> str:
    if check.messages:
        return "; ".join(check.messages)
    return ", ".join(check.reason_codes)


def _link(check: EligibilityCheck) -> str:
    return next((link["resource"] for link in check.approval_links if link.get("resource")), "")


class _Context:
    """Everything the sheets need, loaded with a handful of column-only queries (no ORM objects,
    and never the large raw evidence payloads)."""

    def __init__(self, session: Session, job: Job):
        self.job = job
        seller = session.get(SellerAccount, job.seller_account_id)
        self.seller_ref = seller.amazon_seller_id or str(seller.id)
        self.data_source = (
            "MOCK (demo data, not live Amazon results)"
            if job.data_source == DataSource.MOCK
            else "LIVE"
        )
        self.brands = {
            b.id: b
            for b in session.execute(
                select(
                    JobBrand.id,
                    JobBrand.position,
                    JobBrand.display_name,
                    JobBrand.status,
                    JobBrand.brand_status,
                    JobBrand.brand_summary,
                    JobBrand.last_error,
                    JobBrand.finished_at,
                ).where(JobBrand.job_id == job.id)
            )
        }
        self.checks = session.execute(
            select(
                EligibilityCheck.id,
                EligibilityCheck.asin,
                EligibilityCheck.condition,
                EligibilityCheck.marketplace_id,
                EligibilityCheck.state,
                EligibilityCheck.status,
                EligibilityCheck.rule_id,
                EligibilityCheck.rationale,
                EligibilityCheck.reason_codes,
                EligibilityCheck.messages,
                EligibilityCheck.approval_links,
                EligibilityCheck.verification_methods,
                EligibilityCheck.checked_at,
                EligibilityCheck.job_brand_id,
                CatalogCandidate.title,
            )
            .join(CatalogCandidate, CatalogCandidate.id == EligibilityCheck.candidate_id)
            .where(EligibilityCheck.job_id == job.id)
            .order_by(EligibilityCheck.seq)
        ).all()
        self.observations: dict[Any, list[Any]] = {}
        for obs in session.execute(
            select(
                CheckObservation.check_id,
                CheckObservation.evidence_ref,
                CheckObservation.observed_at,
                CheckObservation.outcome,
                CheckObservation.http_status,
            )
            .join(EligibilityCheck, EligibilityCheck.id == CheckObservation.check_id)
            .where(EligibilityCheck.job_id == job.id)
            .order_by(CheckObservation.observed_at, CheckObservation.id)
        ):
            self.observations.setdefault(obs.check_id, []).append(obs)
        self._marketplace = self.checks[0].marketplace_id if self.checks else ""

    def brand_result(self, brand) -> str:
        return brand.brand_status.value if brand.brand_status else ""

    def product_row(self, check) -> list[Any]:
        brand = self.brands[check.job_brand_id]
        observations = self.observations.get(check.id)
        messages = check.messages
        return [
            brand.display_name,
            check.asin,
            check.title or "",
            check.marketplace_id,
            check.condition,
            check.status.value if check.status else "PENDING",
            "; ".join(messages) if messages else ", ".join(check.reason_codes),
            next((link["resource"] for link in check.approval_links if link.get("resource")), ""),
            ", ".join(check.verification_methods),
            self.seller_ref,
            observations[-1].evidence_ref if observations else "",
            check.checked_at,
            brand.brand_status.value if brand.brand_status else "",
            check.rule_id or "",
            check.rationale or "",
            self.data_source,
        ]

    def brand_row(self, brand, status: str, rationale: str, rule: str = "") -> list[Any]:
        return [
            brand.display_name,
            "",
            "",
            self._marketplace,
            "",
            status,
            "",
            "",
            "",
            self.seller_ref,
            "",
            brand.finished_at,
            self.brand_result(brand),
            rule,
            rationale,
            self.data_source,
        ]


def build_workbook(session: Session, job: Job) -> bytes:
    ctx = _Context(session, job)
    done = [c for c in ctx.checks if c.state == CheckState.DONE]

    wb = Workbook(write_only=True)
    wb.properties.title = f"Eligibility report: {job.name or job.id}"
    wb.properties.creator = "Amazon Seller Eligibility Checker"
    if job.data_source == DataSource.MOCK:
        wb.properties.subject = "DEMO DATA generated from offline fixtures, not live Amazon results"

    _write(
        wb.create_sheet(SHEET_LISTABLE),
        PRODUCT_COLUMNS,
        (ctx.product_row(c) for c in done if c.status == EligibilityStatus.VERIFIED_LISTABLE),
    )
    _write(
        wb.create_sheet(SHEET_RESTRICTED),
        PRODUCT_COLUMNS,
        (ctx.product_row(c) for c in done if c.status in RESTRICTED),
    )

    def manual_rows():
        for check in done:
            if check.status == EligibilityStatus.NEEDS_MANUAL_CHECK:
                yield ctx.product_row(check)
        for brand in sorted(ctx.brands.values(), key=lambda b: b.position):
            if brand.brand_status in (
                EligibilityStatus.NO_PRODUCT_FOUND,
                EligibilityStatus.NEEDS_MANUAL_CHECK,
            ):
                summary = brand.brand_summary or {}
                rationale = summary.get("rationale", "")
                if brand.status == BrandStatus.FAILED and brand.last_error:
                    rationale = f"{rationale} Last error: {brand.last_error}"
                yield ctx.brand_row(
                    brand, brand.brand_status.value, rationale, summary.get("rule_id", "")
                )

    _write(wb.create_sheet(SHEET_MANUAL), PRODUCT_COLUMNS, manual_rows())

    audit_columns = AUDIT_EXTRA + PRODUCT_COLUMNS

    def audit_rows():
        for brand in sorted(ctx.brands.values(), key=lambda b: b.position):
            summary = brand.brand_summary or {}
            status = brand.brand_status.value if brand.brand_status else brand.status.value
            row = ctx.brand_row(
                brand, status, summary.get("rationale", ""), summary.get("rule_id", "")
            )
            yield ["BRAND_RESULT", "", ""] + row
        for check in ctx.checks:
            observations = ctx.observations.get(check.id, [])
            row = ctx.product_row(check)
            if not observations:
                yield ["CHECK_PENDING", "", ""] + row
            for obs in observations:
                observed = list(row)
                observed[10] = obs.evidence_ref or ""
                observed[11] = obs.observed_at
                yield ["OBSERVATION", obs.outcome.value, obs.http_status or ""] + observed

    _write(wb.create_sheet(SHEET_AUDIT), audit_columns, audit_rows())

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
