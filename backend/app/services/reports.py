"""Excel export, generated from committed database rows only.

Sheets 1-4 are the contract with the client; "Brand Summary" is appended as a fifth because
the brand-level result is what the original request was about. A workbook produced while a job
is still running says so in its banner, and a workbook from mock data says that instead of
pretending to be real.
"""

import io
from collections.abc import Iterable, Sequence
from datetime import UTC, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.domain.classification import RULES_VERSION
from app.enums import (
    BLOCKING_STATUSES,
    BrandStatus,
    CheckState,
    DataSource,
    EligibilityStatus,
    JobStatus,
)
from app.models import (
    CatalogCandidate,
    CheckObservation,
    EligibilityCheck,
    Job,
    JobBrand,
    SellerAccount,
)

SHEET_LISTABLE = "Verified Listable Products"
SHEET_RESTRICTED = "Restricted Products"
SHEET_MANUAL = "Manual Review"
SHEET_AUDIT = "Full Audit History"
SHEET_BRANDS = "Brand Summary"

_MAX_CELL_CHARS = 32_000
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
_BANNER_FONT = Font(bold=True, color="FFFFFF")
_MOCK_FILL = PatternFill("solid", fgColor="B91C1C")
_LIVE_FILL = PatternFill("solid", fgColor="1F2937")
_PARTIAL_FILL = PatternFill("solid", fgColor="B45309")
_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="E5E7EB")

CHECK_COLUMNS = [
    ("Brand", 22),
    ("ASIN", 14),
    ("Product Title", 60),
    ("Category", 24),
    ("Marketplace", 16),
    ("Condition", 14),
    ("Status", 26),
    ("Approval Reason", 46),
    ("Approval Link", 40),
    ("Rationale", 60),
    ("Verification Method", 22),
    ("Seller Account Ref", 20),
    ("Evidence Ref", 22),
    ("Checked At (UTC)", 20),
]
AUDIT_COLUMNS = [
    ("Brand", 22),
    ("ASIN", 14),
    ("Product Title", 50),
    ("Marketplace", 16),
    ("Condition", 14),
    ("Check Status (current)", 26),
    ("Rule", 28),
    ("Observation Source", 22),
    ("Observation Outcome", 22),
    ("HTTP Status", 12),
    ("Observed At (UTC)", 20),
    ("Error Detail", 40),
    ("Reason Codes", 24),
    ("Messages", 46),
    ("Verification Method", 22),
    ("Seller Account Ref", 20),
    ("Evidence Ref", 22),
    ("Rules Version", 14),
]
BRAND_COLUMNS = [
    ("Brand (as supplied)", 24),
    ("Searched As", 22),
    ("Run Status", 14),
    ("Brand-Level Result", 26),
    ("What It Means", 44),
    ("Candidates Found", 12),
    ("Brand-Matched", 12),
    ("ASINs Checked", 12),
    ("Listable ASINs", 36),
    ("Sample Sufficient", 12),
    ("Rationale", 70),
    ("Attempts", 10),
    ("Last Error", 40),
    ("Finished At (UTC)", 20),
]

_BRAND_RESULT_MEANING = {
    EligibilityStatus.VERIFIED_LISTABLE: "Every sampled ASIN verified listable. Says nothing "
    "about ASINs that were not checked.",
    EligibilityStatus.BRAND_APPROVAL_REQUIRED: "Brand approval observed on every sampled ASIN; "
    "none verified listable.",
    EligibilityStatus.OTHER_APPROVAL_REQUIRED: "A non-brand approval (e.g. category) blocks the "
    "sampled ASINs.",
    EligibilityStatus.NOT_ELIGIBLE: "Amazon reports the seller as not eligible for the sampled "
    "ASINs.",
    EligibilityStatus.MIXED_RESULTS: "Checked ASINs disagree; see the product sheets for which "
    "ASINs may be listed.",
    EligibilityStatus.NEEDS_MANUAL_CHECK: "Unverified: not enough reliable evidence to decide.",
    EligibilityStatus.NO_PRODUCT_FOUND: "No catalog product whose brand attribute matches.",
}


# ---- cell helpers --------------------------------------------------------------------------


def safe_text(value: str) -> str:
    """Neutralise spreadsheet formula injection in untrusted text (titles, brand names)."""
    text = value[:_MAX_CELL_CHARS]
    return "'" + text if text.startswith(_FORMULA_PREFIXES) else text


def _naive_utc(value: datetime) -> datetime:
    return value.astimezone(UTC).replace(tzinfo=None)


def _cell(ws, value: Any) -> WriteOnlyCell:
    if isinstance(value, datetime):
        cell = WriteOnlyCell(ws, _naive_utc(value))
        cell.number_format = "yyyy-mm-dd hh:mm:ss"
        return cell
    if isinstance(value, str):
        cell = WriteOnlyCell(ws, safe_text(value))
        cell.data_type = "s"  # never let a string be read back as a formula
        return cell
    if value is None:
        return WriteOnlyCell(ws, None)
    if isinstance(value, bool):
        return WriteOnlyCell(ws, "Yes" if value else "No")
    return WriteOnlyCell(ws, value)


def _row(ws, values: Iterable[Any]) -> list[WriteOnlyCell]:
    return [_cell(ws, v) for v in values]


def _banner_cell(ws, text: str, fill: PatternFill) -> WriteOnlyCell:
    cell = WriteOnlyCell(ws, text)
    cell.data_type = "s"
    cell.font = _BANNER_FONT
    cell.fill = fill
    cell.alignment = Alignment(vertical="center")
    return cell


def _start_sheet(wb: Workbook, title: str, columns, banner: str, fill: PatternFill):
    ws = wb.create_sheet(title)
    for index, (_, width) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(columns))}2"
    ws.append([_banner_cell(ws, banner, fill)])
    header = []
    for name, _ in columns:
        cell = WriteOnlyCell(ws, name)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        header.append(cell)
    ws.append(header)
    return ws


# ---- row builders --------------------------------------------------------------------------


def seller_reference(seller: SellerAccount) -> str:
    return seller.amazon_seller_id or f"acct:{str(seller.id)[:8]}"


def _reason_text(check: EligibilityCheck) -> str:
    if check.messages:
        return "; ".join(check.messages)
    return ", ".join(check.reason_codes)


def _first_link(check: EligibilityCheck) -> str:
    for link in check.approval_links or []:
        if link.get("resource"):
            return link["resource"]
    return ""


def _evidence_ref(observation: CheckObservation | None) -> str:
    if observation is None:
        return ""
    return observation.evidence_ref or f"observation:{observation.id}"


def _check_values(
    check: EligibilityCheck,
    brand_name: str,
    candidate: CatalogCandidate,
    seller_ref: str,
    latest: CheckObservation | None,
) -> list[Any]:
    return [
        brand_name,
        check.asin,
        candidate.title,
        candidate.category_name,
        check.marketplace_id,
        check.condition,
        check.status.value if check.status else "NOT CHECKED",
        _reason_text(check),
        _first_link(check),
        check.rationale,
        ", ".join(check.verification_methods),
        seller_ref,
        _evidence_ref(latest),
        check.checked_at,
    ]


def _audit_values(
    check: EligibilityCheck,
    brand_name: str,
    candidate: CatalogCandidate,
    seller_ref: str,
    obs: CheckObservation | None,
) -> list[Any]:
    signals_codes = ""
    messages = ""
    if obs is not None and obs.signals:
        reasons = [
            r
            for entry in obs.signals.get("entries", [])
            for r in entry.get("reasons", [])
            if isinstance(r, dict)
        ]
        signals_codes = ", ".join(str(r["code"]) for r in reasons if r.get("code"))
        messages = "; ".join(str(r["message"]) for r in reasons if r.get("message"))
    return [
        brand_name,
        check.asin,
        candidate.title,
        check.marketplace_id,
        check.condition,
        check.status.value if check.status else "NOT CHECKED",
        check.rule_id,
        obs.source.value if obs else None,
        obs.outcome.value if obs else None,
        obs.http_status if obs else None,
        obs.observed_at if obs else None,
        obs.error_detail if obs else None,
        signals_codes,
        messages,
        ", ".join(check.verification_methods),
        seller_ref,
        _evidence_ref(obs),
        check.rules_version,
    ]


def _brand_values(brand: JobBrand) -> list[Any]:
    summary = brand.brand_summary or {}
    status = brand.brand_status
    return [
        brand.original_name.strip(),
        brand.display_name,
        brand.status.value,
        status.value if status else "NOT FINISHED",
        _BRAND_RESULT_MEANING.get(status, "") if status else "",
        summary.get("candidates_total"),
        summary.get("candidates_brand_matched"),
        summary.get("asins_checked"),
        ", ".join(summary.get("listable_asins", [])),
        summary.get("sample_sufficient"),
        summary.get("rationale"),
        brand.attempts,
        brand.last_error,
        brand.finished_at,
    ]


def _manual_brand_values(brand: JobBrand, seller_ref: str) -> list[Any]:
    """A brand-level row for Manual Review: nothing was checked, so there is no ASIN."""
    summary = brand.brand_summary or {}
    status = brand.brand_status.value if brand.brand_status else "NOT FINISHED"
    reason = brand.last_error or summary.get("rationale") or ""
    return [
        brand.display_name,
        None,
        None,
        None,
        None,
        None,
        status,
        reason,
        None,
        summary.get("rationale"),
        None,
        seller_ref,
        None,
        brand.finished_at,
    ]


# ---- the workbook --------------------------------------------------------------------------


def _banner_for(job: Job, seller_ref: str, generated: datetime) -> tuple[str, PatternFill]:
    stamp = generated.astimezone(UTC).strftime("%Y-%m-%d %H:%M UTC")
    state = job.status.value
    label = job.name or str(job.id)
    if job.data_source == DataSource.MOCK:
        return (
            f"DEMO DATA - offline fixtures, NOT results from Amazon. Says nothing about what "
            f"seller {seller_ref} can list. Job: {label}. Generated {stamp}.",
            _MOCK_FILL,
        )
    partial = job.status != JobStatus.COMPLETED
    text = (
        f"Amazon SP-API data for seller {seller_ref} | Job: {label} | Status: {state}"
        + (" (PARTIAL: some brands are not finished)" if partial else "")
        + f" | Rules {RULES_VERSION} | Generated {stamp}. "
        "Statuses describe the ASINs checked, not whole brands."
    )
    return text, _PARTIAL_FILL if partial else _LIVE_FILL


def _latest_observations(session: Session, job_id) -> dict[Any, CheckObservation]:
    rows = session.scalars(
        select(CheckObservation)
        .join(EligibilityCheck, EligibilityCheck.id == CheckObservation.check_id)
        .where(EligibilityCheck.job_id == job_id)
        .order_by(CheckObservation.id)
    )
    return {obs.check_id: obs for obs in rows}  # later rows overwrite earlier: latest wins


def build_workbook(session: Session, job: Job, *, generated: datetime) -> bytes:
    seller = session.get(SellerAccount, job.seller_account_id)
    seller_ref = seller_reference(seller)
    banner, fill = _banner_for(job, seller_ref, generated)

    wb = Workbook(write_only=True)
    listable = _start_sheet(wb, SHEET_LISTABLE, CHECK_COLUMNS, banner, fill)
    restricted = _start_sheet(wb, SHEET_RESTRICTED, CHECK_COLUMNS, banner, fill)
    manual = _start_sheet(wb, SHEET_MANUAL, CHECK_COLUMNS, banner, fill)
    audit = _start_sheet(wb, SHEET_AUDIT, AUDIT_COLUMNS, banner, fill)
    brands_ws = _start_sheet(wb, SHEET_BRANDS, BRAND_COLUMNS, banner, fill)

    latest = _latest_observations(session, job.id)

    checks = session.execute(
        select(EligibilityCheck, JobBrand.display_name, CatalogCandidate)
        .join(JobBrand, JobBrand.id == EligibilityCheck.job_brand_id)
        .join(CatalogCandidate, CatalogCandidate.id == EligibilityCheck.candidate_id)
        .where(EligibilityCheck.job_id == job.id)
        .order_by(JobBrand.position, EligibilityCheck.asin, EligibilityCheck.condition)
    ).yield_per(1000)
    for check, brand_name, candidate in checks:
        if check.state != CheckState.DONE or check.status is None:
            continue
        values = _check_values(check, brand_name, candidate, seller_ref, latest.get(check.id))
        if check.status == EligibilityStatus.VERIFIED_LISTABLE:
            listable.append(_row(listable, values))
        elif check.status in BLOCKING_STATUSES:
            restricted.append(_row(restricted, values))
        else:
            manual.append(_row(manual, values))

    all_brands = session.scalars(
        select(JobBrand).where(JobBrand.job_id == job.id).order_by(JobBrand.position)
    ).all()
    for brand in all_brands:
        brands_ws.append(_row(brands_ws, _brand_values(brand)))
        needs_review = brand.status == BrandStatus.FAILED or brand.brand_status in (
            EligibilityStatus.NO_PRODUCT_FOUND,
            EligibilityStatus.NEEDS_MANUAL_CHECK,
        )
        if needs_review:
            manual.append(_row(manual, _manual_brand_values(brand, seller_ref)))

    history = session.execute(
        select(EligibilityCheck, JobBrand.display_name, CatalogCandidate, CheckObservation)
        .join(JobBrand, JobBrand.id == EligibilityCheck.job_brand_id)
        .join(CatalogCandidate, CatalogCandidate.id == EligibilityCheck.candidate_id)
        .outerjoin(CheckObservation, CheckObservation.check_id == EligibilityCheck.id)
        .where(EligibilityCheck.job_id == job.id)
        .order_by(
            JobBrand.position,
            EligibilityCheck.asin,
            EligibilityCheck.condition,
            CheckObservation.observed_at,
            CheckObservation.id,
        )
    ).yield_per(1000)
    for check, brand_name, candidate, obs in history:
        audit.append(_row(audit, _audit_values(check, brand_name, candidate, seller_ref, obs)))

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def report_filename(job: Job) -> str:
    stem = "".join(c if c.isalnum() or c in "-_" else "_" for c in (job.name or "brand-check"))
    return f"{stem[:60] or 'brand-check'}-{str(job.id)[:8]}.xlsx"


def sheet_names() -> Sequence[str]:
    return (SHEET_LISTABLE, SHEET_RESTRICTED, SHEET_MANUAL, SHEET_AUDIT, SHEET_BRANDS)


__all__ = ["build_workbook", "report_filename", "safe_text", "sheet_names"]
