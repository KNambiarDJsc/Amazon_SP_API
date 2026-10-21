"""Turn pasted text, CSV and XLSX uploads into a validated, de-duplicated brand list.

Everything here is pure (bytes/str in, result out) so it is easy to test and has no
opinion about HTTP or the database. Uploads are untrusted: file-level problems raise
IngestionError; row-level problems are reported per row and never abort the job.
"""

import csv
import io
import zipfile
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field

from openpyxl import load_workbook

from app.config import get_settings
from app.domain.brand import clean_whitespace, dedupe_key, match_key
from app.schemas import DuplicateRow, IngestionReport, InvalidRow

# A spreadsheet row number paired with the raw cell text (None for an empty cell).
Row = tuple[int, str | None]

_HEADER_NAMES = {"brand", "brands", "brand name", "brand_name", "brandname"}
_MAX_UNCOMPRESSED_BYTES = 64 * 1024 * 1024
_MAX_ZIP_ENTRIES = 2000
_ROW_SCAN_SLACK = 1000  # blank/duplicate rows tolerated beyond max_brands_per_job


class IngestionError(ValueError):
    """The input as a whole is unusable. The message is safe to show to the client."""


@dataclass(frozen=True)
class AcceptedBrand:
    position: int  # 1-based order of acceptance
    row: int  # row in the source the brand came from
    original_name: str  # exactly as supplied
    display_name: str  # whitespace-normalised, case kept
    dedupe_key: str


@dataclass
class IngestionResult:
    accepted: list[AcceptedBrand] = field(default_factory=list)
    report: IngestionReport = field(default_factory=IngestionReport)


def _cell_text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, str):
        return value
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # a brand typed as 3 is stored by Excel as 3.0
    return str(value)


def _row_limit() -> int:
    return get_settings().max_brands_per_job + _ROW_SCAN_SLACK


# ---- sources -------------------------------------------------------------------------------


def rows_from_text(text: str) -> list[Row]:
    """One brand per line; line numbers are 1-based."""
    lines = text.lstrip("﻿").splitlines()
    if len(lines) > _row_limit():
        raise IngestionError(f"too many lines (limit {_row_limit()})")
    return [(i, line) for i, line in enumerate(lines, start=1)]


def rows_from_list(values: Sequence[str | None]) -> list[Row]:
    if len(values) > _row_limit():
        raise IngestionError(f"too many rows (limit {_row_limit()})")
    return [(i, v) for i, v in enumerate(values, start=1)]


def _rows_from_table(table: Iterable[Sequence[object]]) -> list[Row]:
    """Pick the brand column: the one headed 'Brand'/'Brand Name', otherwise the first column.

    A header row is only skipped when it is recognised, so a sheet with no header at all
    keeps its first brand.
    """
    limit = _row_limit()
    cells: list[list[str | None]] = []
    for index, raw_row in enumerate(table):
        if index >= limit:
            raise IngestionError(f"too many rows (limit {limit})")
        cells.append([_cell_text(c) for c in raw_row])
    if not cells:
        return []

    column, first_data = 0, 0
    for idx, header in enumerate(cells[0]):
        if header is not None and clean_whitespace(header).casefold() in _HEADER_NAMES:
            column, first_data = idx, 1
            break

    rows: list[Row] = []
    for index in range(first_data, len(cells)):
        row = cells[index]
        rows.append((index + 1, row[column] if column < len(row) else None))
    return rows


def rows_from_csv(data: bytes) -> list[Row]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252", errors="replace")  # Excel on Windows saves legacy CSVs so
    try:
        return _rows_from_table(csv.reader(io.StringIO(text, newline="")))
    except csv.Error as exc:
        raise IngestionError(f"could not parse CSV: {exc}") from exc


def _check_zip(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
    except zipfile.BadZipFile as exc:
        raise IngestionError("not a valid .xlsx file") from exc
    uncompressed = sum(e.file_size for e in entries)
    if len(entries) > _MAX_ZIP_ENTRIES or uncompressed > _MAX_UNCOMPRESSED_BYTES:
        raise IngestionError("workbook is too large once decompressed")


def rows_from_xlsx(data: bytes) -> list[Row]:
    _check_zip(data)
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises a wide variety of types on corrupt files
        raise IngestionError("could not read the workbook") from exc
    try:
        sheet = workbook.active
        if sheet is None:
            return []
        return _rows_from_table(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()


def rows_from_upload(filename: str, data: bytes) -> list[Row]:
    name = filename.lower()
    if name.endswith(".xlsx"):
        return rows_from_xlsx(data)
    if name.endswith(".csv") or name.endswith(".txt"):
        return rows_from_csv(data)
    raise IngestionError("unsupported file type; upload .xlsx, .csv or .txt")


# ---- validation ----------------------------------------------------------------------------


def ingest_brands(rows: Sequence[Row]) -> IngestionResult:
    """Normalise whitespace, drop empties, and collapse case-insensitive duplicates.

    The first occurrence of a brand wins and keeps its original spelling. Nothing is
    dropped silently: every rejected row appears in the report with its row number.
    """
    settings = get_settings()
    result = IngestionResult(report=IngestionReport(total_rows=len(rows)))
    seen: dict[str, int] = {}

    for row_number, raw in rows:
        display = clean_whitespace(raw) if raw is not None else ""
        if not display:
            result.report.empty_rows.append(row_number)
            continue
        if len(display) > settings.max_brand_length:
            result.report.invalid.append(
                InvalidRow(
                    row=row_number,
                    original_name=display[:80] + "…",
                    reason=f"longer than {settings.max_brand_length} characters",
                )
            )
            continue
        if not match_key(display):
            result.report.invalid.append(
                InvalidRow(
                    row=row_number,
                    original_name=display,
                    reason="contains no letters or digits",
                )
            )
            continue

        key = dedupe_key(display)
        if key in seen:
            result.report.duplicates.append(
                DuplicateRow(row=row_number, original_name=raw or "", duplicate_of_row=seen[key])
            )
            continue
        seen[key] = row_number
        result.accepted.append(
            AcceptedBrand(
                position=len(result.accepted) + 1,
                row=row_number,
                original_name=(raw or "")[:500],
                display_name=display,
                dedupe_key=key,
            )
        )

    if len(result.accepted) > settings.max_brands_per_job:
        raise IngestionError(
            f"{len(result.accepted)} brands exceeds the limit of {settings.max_brands_per_job} "
            "per job; split the list"
        )
    result.report.accepted = len(result.accepted)
    return result
